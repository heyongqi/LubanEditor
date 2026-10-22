"""生成可直接交给策划填写的数据表模板（.xlsx）。

关键点：
- A1 必须是 ##var，否则整张 sheet 会被 Luban 忽略（新手第一大坑）；
- 每个字段的「配置方式」决定这一列怎么展开（流式 / sep / 列限定 / 多行）；
- 带 ##group 行与 ## 中文注释行，另附示例数据与若干空行。
"""
from __future__ import annotations

import os
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from .model import ProjectModel, TableDef, FieldDef
from .typecatalog import (parse_type, resolve_kind, _sample_value, _BASE_MAP,
                          elem_of, _bean_fields, _bean_sep, expand_cells,
                          join_format)

HEADER_FILL = PatternFill("solid", fgColor="D9E7F7")
MARK_FILL = PatternFill("solid", fgColor="FBE3D6")
SAMPLE_FILL = PatternFill("solid", fgColor="FFF6DC")

# 多列字段（流式 / 多级表头 / 多行）表头的填充色，由浅到深。
# 连续几个字段都是多列时按顺序取用：相邻两组颜色一定不同，
# 否则一眼看过去分不清「哪几列是同一个字段」。
MULTI_FILLS = ("FFF3B0", "FFE380", "FFD451", "FFC107")


def derive_value_type(full_name: str) -> str:
    """TbItem -> Item；item.TbItem -> item.Item"""
    if not full_name:
        return ""
    mod, _, base = full_name.rpartition(".")
    if base.startswith("Tb") and len(base) > 2:
        base = base[2:]
    return f"{mod}.{base}" if mod else base


class Col:
    """模板里的一列。"""

    def __init__(self, top="", sub="", type="", comment="", group="",
                 sample="", multi: Optional[List[str]] = None):
        self.top = top
        self.sub = sub
        self.type = type
        self.comment = comment
        self.group = group
        self.sample = sample
        self.multi = multi or []      # 非空 = 多行字段，每个元素一行

    @property
    def multiline(self) -> bool:
        return bool(self.multi)


def auto_style(model: ProjectModel, type_str: str) -> str:
    """没显式选配置方式时的默认：复合类型优先用单格 sep（最稳、最省列）。"""
    p = parse_type(type_str)
    if p["kind"] == "container":
        return "sep"
    k = resolve_kind(model, p["base"])
    if k == "bean":
        return "sep"
    return "plain"


def _effective_style(model: ProjectModel, f: FieldDef, default: str) -> str:
    s = (f.fill_style or "auto").strip()
    if s.startswith("tag:"):        # 「引用校验 / 范围校验 / 长度约束」等标签类配置方式
        s = "auto"                  # 不决定列怎么展开，按默认展开
    if s and s != "auto":
        return s
    return default if default != "auto" else auto_style(model, f.type)


def _scalar_sample(model: ProjectModel, f: FieldDef) -> str:
    return _sample_value(f, model)


def _map_pair_sample(model: ProjectModel, elem: str, value: str) -> str:
    """map 的一个 key-value 对示例：key|value（value 是结构时按其 sep 展开）。"""
    k = _sample_value(FieldDef(name="k", type=elem or "int"), model)
    v = _sample_value(FieldDef(name="v", type=value or "string"), model)
    return f"{k}|{v}"


# ---------------------------------------------------------------- 单元格格式
# 官方「紧凑格式」：format 写在数据表 ##var 的**字段名**上（`pos#format=lite`），
# 只有结构 / 容器这类非原子字段才有意义。选 lite/json/lua 后整格就写一个值，
# 不再走「流式 / sep / 分列」那套。下面按官方示例的写法生成示例单元格。
import re as _re


def _lit(s: str, fmt: str) -> str:
    """把示例值写成 json / lua / lite 的字面量。

    lite 是 Luban 自己的紧凑写法，字符串**不加引号**（官方示例 `{Circle,1.5}`）；
    json / lua 才需要引号。
    """
    t = str(s)
    if fmt == "lite":
        return t
    if t in ("true", "false", "null") or _re.fullmatch(r"-?\d+(\.\d+)?", t):
        return t
    return '"%s"' % t.replace('"', '\\"')


def _bean_literal(model: ProjectModel, base: str, fmt: str) -> str:
    fs = _bean_fields(model, base)
    if not fs:
        return "{}"
    vals = [_sample_value(x, model) for x in fs]
    if fmt == "lite":                     # {1,2,3}
        return "{" + ",".join(vals) + "}"
    if fmt == "json":                     # {"x":1,"y":2}
        return "{" + ",".join('"%s":%s' % (x.name, _lit(v, fmt))
                              for x, v in zip(fs, vals)) + "}"
    return "{" + ",".join("%s=%s" % (x.name, _lit(v, fmt))
                          for x, v in zip(fs, vals)) + "}"      # lua: {x=1,y=2}


def _map_keys(model: ProjectModel, ktype: str) -> List[str]:
    k = (ktype or "int").strip()
    if resolve_kind(model, k) == "enum":
        e = model.find_enum(k)
        names = [i.name for i in (e.items if e else [])] or ["A", "B"]
        return (names * 2)[:2]
    if k in ("string", "text"):
        return ["k1", "k2"]
    return ["1", "2"]


def _cell_literal(model: ProjectModel, f: FieldDef, fmt: str) -> str:
    """整格一个值的示例：按 f 的类型生成 lite / json / lua 写法。"""
    p = parse_type(f.type)
    if p["kind"] == "container":
        cont = p["container"]
        elem = p["elem"] or "int"
        ebean = resolve_kind(model, elem) == "bean"
        ev = (lambda: _bean_literal(model, elem, fmt)) if ebean else \
             (lambda: _lit(_sample_value(FieldDef(name="e", type=elem), model), fmt))
        if cont == "map":
            vt = p["value"] or "string"
            vbean = resolve_kind(model, vt) == "bean"
            vv = (lambda: _bean_literal(model, vt, fmt)) if vbean else \
                 (lambda: _lit(_sample_value(FieldDef(name="v", type=vt), model), fmt))
            ks = _map_keys(model, elem)
            if fmt == "json":             # map 的每一项是 [key, value] 两元组
                return "[" + ",".join("[%s,%s]" % (_lit(k, fmt), vv())
                                      for k in ks) + "]"
            if fmt == "lua":              # {[k]=v, ...}
                return "{" + ",".join("[%s]=%s" % (_lit(k, fmt), vv())
                                      for k in ks) + "}"
            return "{" + ",".join("{%s,%s}" % (_lit(k, fmt), vv())
                                  for k in ks) + "}"   # lite：每项一个 {k,v}
        body = ",".join(ev() for _ in range(2))
        if fmt == "json":
            return "[" + body + "]"
        return "{" + body + "}"           # lite / lua 都是花括号
    if resolve_kind(model, p["base"]) == "bean":
        return _bean_literal(model, p["base"], fmt)
    return _lit(_sample_value(f, model), fmt)


def build_cols(model: ProjectModel, fields: List[FieldDef], default: str = "sep") -> List[Col]:
    """按每个字段的配置方式展开成实际列。"""
    cols: List[Col] = []
    for f in fields:
        style = _effective_style(model, f, default)
        p = parse_type(f.type)
        base = p["base"]
        grp = f.group or ""
        cmt = f.comment or f.name

        # ---------- 单元格格式（lite / json / lua）----------
        # 选了紧凑格式就整格一个值：字段名带 #format=，不再按配置方式展开列。
        fmt = (getattr(f, "fmt", "") or "").strip().lower()
        if fmt in ("lite", "json", "lua") and (
                p["kind"] == "container" or resolve_kind(model, base) == "bean"):
            cols.append(Col(join_format(f.name, fmt), "", f.type, cmt, grp,
                            _cell_literal(model, f, fmt)))
            continue

        # ---------- 容器 ----------
        if p["kind"] == "container":
            cont = p["container"]
            elem = elem_of(model, f.type)
            ek = resolve_kind(model, elem)
            fs = _bean_fields(model, elem) if ek == "bean" else []
            esep = _bean_sep(model, elem)
            es = _scalar_sample(model, FieldDef(name="e", type=elem))

            if style == "multiline":
                if ek == "bean" and fs and cont != "map":
                    vals = [_scalar_sample(model, x) for x in fs]
                    # 主列承载第一个子字段，其余子字段各占一列；每个元素一行
                    cols.append(Col(f"*{f.name}", fs[0].name, f"{cont},{elem}",
                                    fs[0].comment or fs[0].name, grp,
                                    multi=[vals, vals]))
                    for x in fs[1:]:
                        v = _scalar_sample(model, x)
                        cols.append(Col("", x.name, "", x.comment or x.name, grp,
                                        multi=[[v], [v]]))
                else:
                    # map 的类型必须写全 key / value；每行一对 key-value
                    mt = (f"{cont},{elem},{p['value'] or 'string'}"
                          if cont == "map" else f"{cont},{elem}")
                    kv = _map_pair_sample(model, elem, p["value"] or "string")
                    one = [kv if cont == "map" else es]
                    cols.append(Col(f"*{f.name}", "", mt, cmt, grp,
                                    multi=[one, one]))
                continue

            if style in ("stream", "column") and ek == "bean" and fs:
                n = len(fs)
                total = n * 2
                vals = [_scalar_sample(model, x) for x in fs] * 2
                for i in range(total):
                    subname = fs[i % n].name if style == "column" else ""
                    subcmt = (fs[i % n].comment or fs[i % n].name) if style == "column" else ""
                    cols.append(Col(f.name if i == 0 else "", subname,
                                    f"{cont},{elem}" if i == 0 else "",
                                    cmt if i == 0 else subcmt, grp, vals[i]))
                continue

            if style in ("stream", "column") and cont == "map":
                # 键作子列名：每个 key 占一列，格子里填该 key 的 value
                vtype = p["value"] or "string"
                vs = _scalar_sample(model, FieldDef(name="v", type=vtype))
                cols.append(Col(f.name, "key1", f"map,{elem},{vtype}", cmt, grp, vs))
                cols.append(Col("", "key2", "", cmt, grp, vs))
                continue

            # 默认 / sep：一格装下（保留字段自带 #size / #sep 这类容器标签）
            inline = p["sep"] or ""
            if "#sep=" not in inline:
                inline += "#sep=|"
            if cont == "map":
                # map 要写全 key 和 value；value 是结构时再给它一个 sep
                vtype = p["value"] or "string"
                vwrap = (f"({vtype}#sep={_bean_sep(model, vtype)})"
                         if resolve_kind(model, vtype) == "bean" else vtype)
                t = f"({cont}{inline}),{elem},{vwrap}"
                kv = _map_pair_sample(model, elem, vtype)
                cols.append(Col(f.name, "", t, cmt, grp, "|".join([kv, kv])))
            elif ek == "bean" and fs:
                one = esep.join(_scalar_sample(model, x) for x in fs)
                t = f"({cont}{inline}),({elem}#sep={esep})"
                cols.append(Col(f.name, "", t, cmt, grp, "|".join([one, one])))
            else:
                t = f"({cont}{inline}),{elem}"
                cols.append(Col(f.name, "", t, cmt, grp, "|".join([es, es, es])))
            continue

        # ---------- 结构 bean ----------
        k = resolve_kind(model, base)
        if k == "bean":
            fs = _bean_fields(model, base)
            sep = _bean_sep(model, base)
            vals = [_scalar_sample(model, x) for x in fs]
            subs = model.subclasses_of(base)

            if style == "nullable" or p["nullable"]:
                cols.append(Col(f.name, "", base + "?", cmt, grp, "null"))
                continue
            if style == "poly" or (model.is_polymorphic(base) and style == "sep" and subs):
                sname = subs[0].short_name
                cols.append(Col(f.name, "", f"{base}#sep={sep}", cmt, grp,
                                sep.join([sname] + vals)))
                continue
            if style == "column" and fs:
                for i, x in enumerate(fs):
                    cols.append(Col(f.name if i == 0 else "", x.name,
                                    base if i == 0 else "",
                                    cmt if i == 0 else (x.comment or x.name), grp, vals[i]))
                continue
            if style == "stream" and fs:
                for i, x in enumerate(fs):
                    cols.append(Col(f.name if i == 0 else "", "",
                                    base if i == 0 else "",
                                    cmt if i == 0 else "", grp, vals[i]))
                continue
            # sep（默认）
            if fs:
                cols.append(Col(f.name, "", f"{base}#sep={sep}", cmt, grp, sep.join(vals)))
            else:
                cols.append(Col(f.name, "", base, cmt, grp, "{}"))
            continue

        # ---------- 枚举 ----------
        if k == "enum":
            e = model.find_enum(base)
            if style == "nullable":
                cols.append(Col(f.name, "", base + "?", cmt, grp, "null"))
                continue
            names = [i.name for i in (e.items if e else [])] or ["A", "B"]
            if style == "flags" and len(names) >= 2:
                cols.append(Col(f.name, "", base, cmt, grp, "|".join(names[:2])))
            else:
                cols.append(Col(f.name, "", f.type, cmt, grp, names[0]))
            continue

        # ---------- 基础类型 ----------
        if style == "nullable":
            cols.append(Col(f.name, "", base + "?", cmt, grp, "null"))
        elif style == "bang":
            cols.append(Col(f.name, "", base + "!", cmt, grp, _scalar_sample(model, f)))
        else:
            cols.append(Col(f.name, "", f.type, cmt, grp, _scalar_sample(model, f)))
    return cols


def _style_header(ws, row: int, ncols: int, mark: bool = False):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True)
        cell.fill = MARK_FILL if (mark and c == 1) else HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def col_groups(cols: List[Col]) -> List[Tuple[int, int]]:
    """(起始列下标, 跨列数)：top 为空的列是左边那个字段的续列。

    一个字段在 Excel 里横跨几列时，Luban 认的是「第一列写字段名，
    后面几列跟着」——所以这里把续列归到同一个组，用来做合并与配色。
    """
    groups: List[List[int]] = []
    for i, c in enumerate(cols):
        if c.top or not groups:
            groups.append([i, 0])
        groups[-1][1] += 1
    return [(g[0], g[1]) for g in groups]


def _column_notes(cols: List[Col], groups: List[Tuple[int, int]]) -> List[str]:
    """多列字段的「这一列到底填什么」备注。

    流式填法下子字段名是空的，光看表头分不清第 3 列是第几个元素；
    这一行写 `items[3/4]`（多级表头有子字段名时写 `items.count`）。
    用 `##` 标记写：Luban 允许有多行 `##`，取第一行当字段注释，
    所以这行只是给人看的，不影响导出；换成 `##comment` 反而会抢走注释。
    """
    out = [""] * len(cols)
    for start, span in groups:
        if span < 2:
            continue
        top = (cols[start].top or "").lstrip("*")
        for k in range(span):
            sub = (cols[start + k].sub or "").strip()
            out[start + k] = f"{top}.{sub}" if sub else f"{top}[{k + 1}/{span}]"
    return out


def _style_at(ws, row: int, col: int, fill, bold: bool = True):
    cell = ws.cell(row=row, column=col)
    cell.font = Font(bold=bold)
    cell.fill = fill
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _fill_of(shade: Dict[int, int], gi: int):
    """第 gi 组表头的填充色：多列组按序号取不同深浅的黄。"""
    s = shade.get(gi)
    if s is None:
        return HEADER_FILL
    return PatternFill("solid", fgColor=MULTI_FILLS[s % len(MULTI_FILLS)])


def _write_merged_row(ws, row: int, marker: str, groups: List[Tuple[int, int]],
                      values: List, shade: Dict[int, int]) -> int:
    """写一行**需要合并**的表头（##var / ##type / ##group）。

    多列字段的续列必须合并 —— 不合并的话，看表的人会以为它们是几个独立字段，
    而且 Luban 认的是「第一列是字段名，右边空列跟着」，分开写反而容易填错。
    """
    ws.cell(row=row, column=1, value=marker)
    _style_at(ws, row, 1, MARK_FILL)
    # 先给整段上色再合并：合并后续格会变成 MergedCell，
    # 有些表格软件渲染合并区时只看左上角，提前上色两种情况都对。
    for gi, (start, span) in enumerate(groups):
        c0 = 2 + start
        ws.cell(row=row, column=c0, value=values[start])
        fill = _fill_of(shade, gi)
        for k in range(span):
            _style_at(ws, row, c0 + k, fill)
    for _gi, (start, span) in enumerate(groups):
        if span > 1:
            c0 = 2 + start
            ws.merge_cells(start_row=row, start_column=c0,
                           end_row=row, end_column=c0 + span - 1)
    return row + 1


def _write_plain_row(ws, row: int, marker: str, values: List) -> int:
    """写一行不合并的表头（子字段名 / 中文说明 / 列序备注）。"""
    ws.cell(row=row, column=1, value=marker)
    _style_at(ws, row, 1, MARK_FILL)
    for i, v in enumerate(values):
        ws.cell(row=row, column=2 + i, value=v)
        _style_at(ws, row, 2 + i, HEADER_FILL)
    return row + 1


def generate_data_xlsx(model: ProjectModel, table: TableDef, fields: List[FieldDef],
                       out_path: str, vertical: bool = False,
                       with_group: bool = True, sample_rows: int = 2,
                       empty_rows: int = 5, style: str = "sep") -> tuple:
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = (table.full_name.split(".")[-1] or "Sheet1")[:31]

    if vertical:
        _write_vertical(ws, model, table, fields)
    else:
        _write_horizontal(ws, model, table, fields, with_group, sample_rows, empty_rows, style)

    try:
        wb.save(out_path)
    except Exception as e:
        return False, f"保存失败：{e}"
    return True, out_path


def _write_horizontal(ws, model, table, fields, with_group, sample_rows, empty_rows, style):
    cols = build_cols(model, fields, default=style)
    groups = col_groups(cols)
    has_sub = any(c.sub for c in cols)
    n = len(cols)
    row = 1

    # 只给跨多列的组编号，用它取不同深浅的黄（相邻两组颜色必然不同）
    shade: Dict[int, int] = {}
    seq = 0
    for gi, (_start, span) in enumerate(groups):
        if span > 1:
            shade[gi] = seq
            seq += 1

    row = _write_merged_row(ws, row, "##var", groups, [c.top for c in cols], shade)
    if has_sub:
        row = _write_plain_row(ws, row, "##var", [c.sub for c in cols])
    row = _write_merged_row(ws, row, "##type", groups, [c.type for c in cols], shade)
    if with_group:
        row = _write_merged_row(ws, row, "##group", groups, [c.group for c in cols], shade)
    row = _write_plain_row(ws, row, "##",
                           [c.comment or c.top or c.sub for c in cols])

    # 多列字段再补一行「每列是什么」，没有多列字段时整行省略（不给普通表添乱）。
    # 注意用 ## 而不是 ##comment：##comment 会被 Luban 当成「代码注释行」，
    # 用它会把字段注释抢成 "items[1/4]"。多个 ## 行是允许的，取第一个当注释。
    notes = _column_notes(cols, groups)
    if any(notes):
        row = _write_plain_row(ws, row, "##", notes)

    ws.freeze_panes = ws.cell(row=row, column=1)

    max_multi = max([len(c.multi) for c in cols if c.multiline] or [0])
    total_rows = max(sample_rows, max_multi)
    for r in range(total_rows):
        for i, c in enumerate(cols):
            if c.multiline:
                v = c.multi[r] if r < len(c.multi) else [""]
                v = v[0] if isinstance(v, list) else v
            else:
                v = c.sample if r == 0 else ""
            cell = ws.cell(row=row, column=2 + i, value=(v if v not in ("", None) else None))
            cell.fill = SAMPLE_FILL
        row += 1

    for _ in range(empty_rows):
        row += 1

    for i, c in enumerate(cols):
        label = c.sub or c.top or ""
        ws.column_dimensions[get_column_letter(2 + i)].width = max(10, min(26, len(str(label)) * 2 + 6))
    ws.column_dimensions["A"].width = 8


def _write_vertical(ws, model, table, fields):
    """纵表：A1=##column，每行一个字段（字段名 | 类型 | 注释 | 值）。"""
    ws.cell(row=1, column=1, value="##column")
    _style_header(ws, 1, 4, mark=True)
    ws.cell(row=2, column=1, value="")
    ws.cell(row=2, column=2, value="##type")
    ws.cell(row=2, column=3, value="##")
    ws.cell(row=2, column=4, value="")
    _style_header(ws, 2, 4)
    r = 3
    for f in fields:
        cols = build_cols(model, [f], default="sep")
        c = cols[0] if cols else Col(f.name, "", f.type)
        ws.cell(row=r, column=1, value=f.name)
        ws.cell(row=r, column=2, value=c.type or f.type)
        ws.cell(row=r, column=3, value=f.comment or f.name)
        ws.cell(row=r, column=4, value=c.sample)
        r += 1
    for cc, w in zip("ABCD", (22, 18, 24, 24)):
        ws.column_dimensions[cc].width = w
    ws.freeze_panes = "A3"


def header_grid(cols: List["Col"], samples: int = 2) -> List[List]:
    """把列列表渲染成预览表头（给界面看，不是真文件）。

    复合字段的续列 top / type 都是空的 —— 它们在数据表里属于左边那个字段，
    不是独立列。这里按「top 为空 = 上一个字段的续列」分组，续列位置放
    `typecatalog.PAD`，界面就会把它们画成一个跨列的合并格，不会看起来像
    一堆莫名的空列。子字段行与注释行仍按列显示（那里本来就有各自的内容）。
    """
    groups = col_groups(cols)

    def merged(vals: List) -> List:
        return [(vals[g[0]], g[1]) for g in groups]

    grid: List[List] = [["##var"] + merged([c.top for c in cols])]
    if any(c.sub for c in cols):
        grid.append(["##var"] + [c.sub for c in cols])
    grid.append(["##type"] + merged([c.type for c in cols]))
    grid.append(["##group"] + merged([c.group for c in cols]))
    grid.append(["##"] + [c.comment or c.top or c.sub for c in cols])
    notes = _column_notes(cols, groups)
    if any(notes):
        grid.append(["##"] + notes)
    for i in range(max(0, samples)):
        grid.append([""] + [col_sample(c) if i == 0 else "" for c in cols])
    return [expand_cells(r) for r in grid]


def col_sample(c: "Col") -> str:
    """预览用的示例值：多行字段的值存在 multi 里，sample 是空的。"""
    if c.sample:
        return c.sample
    if c.multi:
        v = c.multi[0]
        if isinstance(v, list):
            return v[0] if v else ""
        return str(v)
    return ""


def default_input_path(model: ProjectModel, table: TableDef) -> str:
    base = (table.input or "").strip()
    if not base:
        base = f"{table.full_name.split('.')[-1]}.xlsx"
    return os.path.normpath(os.path.join(model.data_dir, base))
