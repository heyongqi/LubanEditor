# -*- coding: utf-8 -*-
"""端到端回归：建工程 → 建结构 → 用界面新建「带容器字段」的表 → 生成模板 → 落盘 → 重开。

和 smoke.py 的分工：
- smoke.py 用**已有**工程跑单点功能（合成工程，不落盘也能测）；
- 本脚本专门盯**容器类型串的完整往返** —— 子类型下拉选出来的 `list,Reward` /
  `map,string,Reward` 必须能写进 `__tables__.xlsx`、生成正确的 `##type` 表头、
  重开工程后子类型下拉还能正确回显。

这条链路曾经有过真 BUG：类型串写成 `(map#sep=|),string,(Reward#sep=,)` 后，
解析层用裸 `split(",")` 把它拆成 `(Reward#sep=` + `)`，重开时子类型只剩半个。
单点测试抓不到这种「落盘再读回」才暴露的问题，所以必须有这一层。

用法：python tests/e2e.py
"""
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ["QT_QPA_PLATFORM"] = "offscreen"

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QMessageBox

from lubankit import ui, exceldef, projectwizard, typecatalog
from lubankit.model import BeanDef, FieldDef

# 临时工程放系统临时目录，不往仓库里丢产物
ROOT = os.path.join(tempfile.gettempdir(), "lubankit_e2e")
FAIL = []


def check(cond, msg):
    print(("  OK   " if cond else "  FAIL ") + msg)
    if not cond:
        FAIL.append(msg)


def main():
    shutil.rmtree(ROOT, ignore_errors=True)

    app = QApplication.instance() or QApplication(sys.argv)
    app.setStyleSheet(ui.QSS)

    # 一切弹窗自动回答，避免脚本卡住
    QMessageBox.information = staticmethod(lambda *a, **k: QMessageBox.Ok)
    QMessageBox.warning = staticmethod(lambda *a, **k: QMessageBox.Ok)
    QMessageBox.critical = staticmethod(lambda *a, **k: QMessageBox.Ok)
    QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.Yes)

    print("== 1. 用工具建一个全新工程，并写入一个结构 Reward ==")
    ok, msg = projectwizard.create_project(ROOT, ["xlsx"])
    check(ok, "创建工程：" + str(msg).splitlines()[0])

    w = ui.MainWindow()
    w.open_project(ROOT)
    check(w.model.initialized, "工程已初始化")

    b = BeanDef(full_name="Reward", sep=",", comment="奖励")
    b.fields = [FieldDef("item_id", "int", "", "道具ID"),
                FieldDef("count", "int", "", "数量")]
    ok, msg = exceldef.save_bean(w.model, b)
    check(ok, "写入结构 Reward：" + str(msg).splitlines()[0])
    exceldef.load_project(w.model)
    w._refresh_tree()
    check(w.model.find_bean("Reward") is not None, "结构写盘后能读回")

    print("== 2. 界面新建表：id(int) + items(list, Reward) + names(map, string, Reward) ==")
    w._open_new_table(ask_leave=False)
    w.nt_input.setText("drop")
    w.nt_name.setText("TbDrop")
    w.nt_mode.set_code("map")
    w.nt_index.setText("id")
    w.nt_fields.item(0, ui.COL_NAME).setText("id")

    w._nt_add_field()
    w.nt_fields.item(1, ui.COL_NAME).setText("items")
    tc = w.nt_fields.cellWidget(1, ui.COL_TYPE)
    names = [tc.itemText(i) for i in range(tc.count())]
    tc.setCurrentIndex([i for i, n in enumerate(names) if n.startswith("list,int")][0])
    subs = w.nt_fields.cellWidget(1, ui.COL_SUB).findChildren(ui.ValueCombo)
    check(len(subs) == 1, "list 只给一个「元素」子类型下拉")
    check([l.text() for l in w.nt_fields.cellWidget(1, ui.COL_SUB)
           .findChildren(ui.QLabel)] == ["元素"], "list 的子类型标「元素」")
    subs[0].set_code("Reward")
    check(w._nt_buf[1].type == "list,Reward",
          "选元素后缓冲区类型串 = %s" % w._nt_buf[1].type)

    w._nt_add_field()
    w.nt_fields.item(2, ui.COL_NAME).setText("names")
    tc2 = w.nt_fields.cellWidget(2, ui.COL_TYPE)
    names2 = [tc2.itemText(i) for i in range(tc2.count())]
    tc2.setCurrentIndex([i for i, n in enumerate(names2)
                         if n.startswith("map,int,string")][0])
    subs2 = w.nt_fields.cellWidget(2, ui.COL_SUB).findChildren(ui.ValueCombo)
    check(len(subs2) == 2, "map 给 key + value 两个子类型下拉")
    check([l.text() for l in w.nt_fields.cellWidget(2, ui.COL_SUB)
           .findChildren(ui.QLabel)] == ["键", "值"], "map 的子类型标「键 / 值」")
    subs2[0].set_code("string")
    subs2[1].set_code("Reward")
    check(w._nt_buf[2].type == "map,string,Reward",
          "map 选完键值后类型串 = %s" % w._nt_buf[2].type)

    # 容器的「配置方式」：默认落在「单单元格 + sep（最省列）」。注意 sep 只是
    # **列展开方式**，不改写类型串本身（①/② 共用同一份类型串），由 template 落成
    # (map#sep=|),K,V；切换配置方式不应该把类型串弄坏。
    subm = w.nt_fields.cellWidget(2, ui.COL_STYLE)
    check(subm.currentData() == "sep",
          "map 字段默认落在「单单元格 + sep」：%s" % subm.currentData())
    subm.setCurrentIndex(subm.findData("stream"))
    check(w._nt_buf[2].fill_style == "stream" and w._nt_buf[2].type == "map,string,Reward",
          "改选「流式」后配置方式 = %s，类型串不受影响：%s"
          % (w._nt_buf[2].fill_style, w._nt_buf[2].type))
    subm.setCurrentIndex(subm.findData("sep"))     # 改回最省列，后面的模板按这个生成
    check(w._nt_buf[2].fill_style == "sep" and w._nt_buf[2].type == "map,string,Reward",
          "改回「单单元格 + sep」后类型串仍完好：%s" % w._nt_buf[2].type)

    check(typecatalog.validate_type(w.model, w._nt_buf[1].type) == [],
          "list 字段类型串无校验问题")
    check(typecatalog.validate_type(w.model, w._nt_buf[2].type) == [],
          "map 字段类型串无校验问题")

    print("== 3. 建表 + 生成模板 + 保存定义 ==")
    w._nt_create()
    w._save_editor()                      # 点底部「保存」才写回 __tables__.xlsx
    fields = [(f.name, f.type) for f in w.fields.get("TbDrop", [])]
    check([n for n, _ in fields] == ["id", "items", "names"],
          "建表后字段齐全：%s" % [n for n, _ in fields])

    data = os.path.join(w.model.data_dir, "drop.xlsx")
    check(os.path.exists(data), "数据表模板已生成：%s" % os.path.basename(data))
    if os.path.exists(data):
        ws = load_workbook(data).worksheets[0]
        row = {ws.cell(row=r, column=1).value: [ws.cell(row=r, column=c).value
                                                for c in range(2, 5)]
               for r in range(1, 6)}
        check(row.get("##var", [])[:3] == ["id", "items", "names"],
              "##var 行：%s" % row.get("##var"))
        t_items = str(row.get("##type", [])[1] or "")
        t_names = str(row.get("##type", [])[2] or "")
        check("list" in t_items and "Reward" in t_items,
              "list 字段的 ##type 含结构名：%s" % t_items)
        check("map" in t_names and "string" in t_names and "Reward" in t_names,
              "map 字段的 ##type 完整（K/V 都在）：%s" % t_names)

    print("== 4. 重开工程：容器类型串必须完整回显 ==")
    w2 = ui.MainWindow()
    w2.open_project(ROOT)
    t = w2.model.find_table("TbDrop")
    check(t is not None, "重开后能读到 TbDrop")
    check(len(w2.fields.get("TbDrop", [])) == 3, "重开后字段数 = 3")

    w2._load_table_editor("TbDrop")
    # 元素类型自己也可以带标签：(Reward#sep=,) 表示「Reward 子字段在一个格子里
    # 用 , 分隔」。这层信息必须原样留着，不能被压成裸的 Reward，否则重开再保存就丢了。
    p1 = typecatalog.parse_type(w2._edit_fields[1].type)
    check(p1["kind"] == "container" and p1["container"] == "list"
          and typecatalog.parse_type(p1["elem"])["base"] == "Reward",
          "重开后 list 行元素类型完整：%s -> elem=%s" % (w2._edit_fields[1].type, p1["elem"]))
    check(p1["sep"] == "#sep=|", "容器自身的 sep 没被吃掉：%s" % p1["sep"])
    p2 = typecatalog.parse_type(w2._edit_fields[2].type)
    check(p2["kind"] == "container" and p2["container"] == "map"
          and p2["elem"] == "string"
          and typecatalog.parse_type(p2["value"])["base"] == "Reward",
          "重开后 map 行的 K/V 完整：%s -> %s / %s"
          % (w2._edit_fields[2].type, p2["elem"], p2["value"]))

    s1 = w2.tbl_fields.cellWidget(1, ui.COL_SUB).findChildren(ui.ValueCombo)
    check(len(s1) == 1 and typecatalog.parse_type(s1[0].code())["base"] == "Reward",
          "重开后第 1 行子类型下拉回显 = %s" % [c.code() for c in s1])
    s2 = w2.tbl_fields.cellWidget(2, ui.COL_SUB).findChildren(ui.ValueCombo)
    check([c.code() for c in s2] == ["string", "(Reward#sep=,)"],
          "重开后第 2 行子类型下拉回显 = %s" % [c.code() for c in s2])
    check("Reward（结构" in s2[1].currentText() and "#" not in s2[1].currentText(),
          "带标签的元素类型显示成可读名：%s" % s2[1].currentText())

    # 「类型」框显示的是带中文名的形式（原始类型串对新手不友好）
    t1 = w2.tbl_fields.cellWidget(1, ui.COL_TYPE).currentText()
    check("（列表）" in t1 and "list" in t1,
          "重开后类型框显示带中文名的形式：%s" % t1)

    print("== 4b. 单元格格式落盘：##var 字段名带 #format=，重开能回显 ==")
    fcb = w2.tbl_fields.cellWidget(1, ui.COL_FMT).findChild(ui.ValueCombo)
    check(fcb is not None, "list 行可以选单元格格式")
    fcb.set_code("json")
    check(w2._edit_fields[1].fmt == "json",
          "选 json 写进缓冲：%s" % w2._edit_fields[1].fmt)
    w2._save_editor()
    check("已保存" in getattr(w2, "_toast", None).text() if getattr(w2, "_toast", None) else False,
          "保存成功后浮出成功提示：%s"
          % (getattr(w2, "_toast", None).text() if getattr(w2, "_toast", None) else "无"))
    data2 = os.path.join(w2.model.data_dir, "drop.xlsx")
    ws2 = load_workbook(data2).worksheets[0]
    hdr = [ws2.cell(row=1, column=c).value for c in range(2, 5)]
    check(hdr[1] == "items#format=json",
          "数据表 ##var 上写了 #format=json：%s" % hdr)
    check(hdr[2] == "names", "没选格式的字段保持原名：%s" % hdr[2])

    w3 = ui.MainWindow()
    w3.open_project(ROOT)
    check(w3.fields.get("TbDrop", [{}])[1].fmt == "json",
          "重开后字段的格式读回来了：%s"
          % (w3.fields.get("TbDrop", [{}])[1].fmt if len(w3.fields.get("TbDrop", [])) > 1 else None))
    w3._load_table_editor("TbDrop")
    check(w3._edit_fields[1].name == "items",
          "字段名没把 #format 带进来：%s" % w3._edit_fields[1].name)
    f3 = w3.tbl_fields.cellWidget(1, ui.COL_FMT).findChild(ui.ValueCombo)
    check(f3 is not None and f3.code() == "json",
          "重开后「单元格格式」回显 = %s" % (f3.code() if f3 else None))

    print("== 5. 清理 ==")
    shutil.rmtree(ROOT, ignore_errors=True)
    check(not os.path.exists(ROOT), "临时工程已删除")


if __name__ == "__main__":
    main()
    print()
    if FAIL:
        print("FAILED: %d" % len(FAIL))
        for m in FAIL:
            print(" - " + m)
        sys.exit(1)
    print("E2E ALL PASSED")
