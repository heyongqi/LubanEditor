# -*- coding: utf-8 -*-
"""冒烟测试：合成工程跑通「创建 → 导入 → 建表 → 保存 → 模板 → 自检 → 界面」。"""
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook, Workbook

from lubankit import (projectconf, exceldef, xmldef, typecatalog, template,
                      runner, knowledge, projectwizard, publish, ui)
from lubankit.model import ProjectModel, TableDef, FieldDef, BeanDef, EnumDef, EnumItem

# 临时工程放系统临时目录：不往仓库里丢测试产物，也避免上一次的残留影响断言
ROOT = os.path.join(tempfile.gettempdir(), "lubankit_smoke")
FAIL = []


def check(cond, msg):
    print(("  OK   " if cond else "  FAIL ") + msg)
    if not cond:
        FAIL.append(msg)


def main():
    print("== 1. 创建工程（xlsx + xml）==")
    shutil.rmtree(ROOT, ignore_errors=True)
    ok, msg = projectwizard.create_project(ROOT, ["xlsx", "xml"])
    check(ok, "创建工程：" + str(msg).splitlines()[0])
    for f in ("luban.conf", "gen.bat", "gen.sh",
              os.path.join("Schemas", "__tables__.xlsx"),
              os.path.join("Data", ""),
              os.path.join("Defines", "tables.xml")):
        check(os.path.exists(os.path.join(ROOT, f)), "生成文件：" + f)
    import json
    conf = json.load(open(os.path.join(ROOT, "luban.conf"), encoding="utf-8"))
    _schemas = [s["fileName"] for s in conf["schemaFiles"]]
    check(any(s["type"] == "table" for s in conf["schemaFiles"]), "luban.conf 登记了表定义")
    check(any(s.get("fileName") == "Defines" for s in conf["schemaFiles"]),
          "luban.conf 登记了 Defines")
    check("Schemas/__tables__.xlsx" in _schemas,
          "表声明 xlsx 放在 Schemas 下：%s" % _schemas)
    check("Schemas/__beans__.xlsx" in _schemas and "Schemas/__enums__.xlsx" in _schemas,
          "生成了 __beans__.xlsx / __enums__.xlsx 并登记：%s" % _schemas)
    check(conf.get("dataDir") == "Data", "数据表目录仍是 Data：%s" % conf.get("dataDir"))

    print("== 2. 任意目录都能打开（无要求）==")
    empty = os.path.join(ROOT, "随便一个空目录")
    os.makedirs(empty, exist_ok=True)
    me = projectconf.discover(empty)
    check(True, "空目录不抛异常")
    check(not me.initialized, "空目录被识别为「未初始化」")
    check(any("初始化" in w for w in me.warnings), "给出初始化提示而非报错")

    print("== 3. 导入刚创建的工程 ==")
    m = projectconf.discover(ROOT)
    exceldef.load_project(m)
    check(m.initialized, "识别为已初始化工程")
    check(len(m.tables) == 0 and len(m.beans) == 0, "新工程定义为空表")
    check(m.groups == ["c", "s", "e"], "分组解析：%s" % m.groups)

    print("== 4. 新建枚举 / 结构（自定义类型）==")
    e = EnumDef(full_name="Quality", comment="品质")
    e.items = [EnumItem("RED", "红", "1", ""), EnumItem("BLUE", "蓝", "2", "")]
    ok, msg = exceldef.save_enum(m, e)
    check(ok, "写入枚举：" + str(msg)[:40])
    b = BeanDef(full_name="Reward", sep=",", comment="奖励")
    b.fields = [FieldDef("item_id", "int", "", "道具ID"),
                FieldDef("count", "int", "", "数量")]
    ok, msg = exceldef.save_bean(m, b)
    check(ok, "写入结构：" + str(msg)[:40])
    m2 = projectconf.discover(ROOT)
    exceldef.load_project(m2)
    check(m2.find_enum("Quality") and len(m2.find_enum("Quality").items) == 2,
          "重新导入能看到枚举项")
    check(m2.find_bean("Reward") and len(m2.find_bean("Reward").fields) == 2,
          "重新导入能看到结构字段")
    check(m2.find_bean("Reward").scope == "common", "无前缀的类型归为公共")

    print("== 4b. 表定义：写入 -> 重新读取（往返）==")
    # 建工程模板会在表头下面写一行「## + 中文列名」的行说明。它不能被当成
    # 二级表头 —— 一旦被当二级表头，每一列都带上 sub，表既读不出来也写不进去，
    # 而且以前还会报告「已保存」。所以这里必须真的落盘再读回来。
    t_rt = TableDef(full_name="TbRound", value_type="Reward",
                    read_schema_from_file="false", input="round.xlsx",
                    index="id", mode="map", group="c,s", comment="往返表")
    m.tables.append(t_rt)
    ok, msg = exceldef.save_tables(m)
    check(ok, "写入表定义：" + str(msg)[:40])
    check(load_workbook(m.table_file).worksheets[0].max_row >= 3, "表定义真的落到了文件里")
    m3 = projectconf.discover(ROOT)
    exceldef.load_project(m3)
    got = m3.find_table("TbRound")
    check(got is not None, "重新导入能看到表：%s" % [t.full_name for t in m3.tables])
    if got:
        check((got.input, got.mode, got.index, got.group, got.comment)
              == ("round.xlsx", "map", "id", "c,s", "往返表"),
              "表定义各列往返一致：%s"
              % [got.input, got.mode, got.index, got.group, got.comment])
    check("一行表都没解析出来" not in " ".join(m3.warnings),
          "没有「一行表都没解析出来」的告警：%s" % m3.warnings)
    _ds, _wb, _w = exceldef.open_def_sheet(m.table_file, exceldef.TABLE_ALIASES)
    check(_ds is not None and all(not c["sub"] for c in _ds.columns),
          "「##」行说明没被当成二级表头：%s"
          % [(c["l0"], c["sub"]) for c in (_ds.columns if _ds else [])][:3])
    if _wb is not None:
        _wb.close()

    print("== 5. 标签多选：不重复追加 ==")
    core, tags = typecatalog.split_tags("int#ref=TbItem!")
    check(core == "int" and "!" in tags, "拆解类型与标签：%s / %s" % (core, tags))
    t1 = typecatalog.join_type("int", ["#ref=TbItem"], nullable=False)
    t2 = typecatalog.join_type(t1, ["#ref=TbItem"], nullable=False)
    check("#ref=TbItem" in t2 and t2.count("#ref=") == 1, "重复应用同一标签不会叠加：%s" % t2)
    t3 = typecatalog.join_type("int", ["!"], nullable=True)
    check(t3.endswith("?") and t3.count("!") == 1, "可空与禁止默认值可共存：%s" % t3)
    t4 = typecatalog.join_type("list,int", ["#sep=|"], nullable=False)
    check(t4.startswith("(list#sep=|)"), "容器标签紧贴容器：%s" % t4)

    print("== 5b. 带括号的元素类型参数不能被逗号拆坏（回归）==")
    # (Reward#sep=,) 里的逗号属于**一个**类型参数。用裸 split(",") 会把它拆成
    # "(Reward#sep=" 和 ")"，落盘再读回时子类型就残缺了 —— 必须括号感知地拆。
    p1 = typecatalog.parse_type("(map#sep=|),string,(Reward#sep=,)")
    check(p1["kind"] == "container" and p1["container"] == "map"
          and p1["sep"] == "#sep=|", "容器的 sep 与参数分开取：%s" % p1["sep"])
    check(p1["elem"] == "string" and p1["value"] == "(Reward#sep=,)",
          "map 的 K / V 完整：%s / %s" % (p1["elem"], p1["value"]))
    p2 = typecatalog.parse_type("(list#sep=|),(Reward#sep=,)")
    check(p2["elem"] == "(Reward#sep=,)", "list 的元素类型保留括号：%s" % p2["elem"])
    check(typecatalog.parse_type(p2["elem"])["base"] == "Reward"
          and typecatalog.parse_type(p2["elem"])["sep"] == "#sep=,",
          "再解析一层能拿回结构与它的 sep")
    core2, _tags2 = typecatalog.split_tags("(map#sep=|),string,(Reward#sep=,)")
    check("(Reward#sep=,)" in core2, "split_tags 不破坏括号参数：%s" % core2)
    t5 = typecatalog.join_type(core2, ["#size=4"], nullable=False)
    check(t5.count("#size=4") == 1 and "(Reward#sep=,)" in t5,
          "加标签后括号参数仍在：%s" % t5)
    check(typecatalog.validate_type(m2, "(map#sep=|),string,(Reward#sep=,)") == [],
          "带括号参数的容器类型串能通过校验")

    print("== 6. 配置方式（随类型变化）==")
    opts = typecatalog.fill_style_options(m2, "list,Reward")
    keys = [k for k, _ in opts]
    check("sep" in keys and "multiline" in keys and "column" in keys,
          "容器给出 sep / 列限定 / 多行：%s" % keys)
    opts2 = typecatalog.fill_style_options(m2, "Reward")
    check(any(k == "column" for k, _ in opts2), "结构给出列限定：%s" % [k for k, _ in opts2])
    opts3 = typecatalog.fill_style_options(m2, "int")
    keys3 = [k for k, _ in opts3]
    check(keys3[:3] == ["plain", "nullable", "bang"],
          "基础类型先给出 直接填/可空/禁止默认值：%s" % keys3)
    check(all("tag:" + k in keys3 for k in ("ref", "range", "path")),
          "基础类型的下拉补齐了类型助手里的校验标签：%s" % keys3)
    opts4 = typecatalog.fill_style_options(m2, "string")
    check("tag:text" in [k for k, _ in opts4], "string 额外给出本地化 #text")
    opts5 = typecatalog.fill_style_options(m2, "Quality")
    check("tag:ref" in [k for k, _ in opts5], "枚举给出引用校验：%s" % [k for k, _ in opts5])
    opts6 = typecatalog.fill_style_options(m2, "list,Reward")
    check("tag:size" in [k for k, _ in opts6], "容器给出长度约束：%s" % [k for k, _ in opts6])
    for opts in (opts3, opts4, opts5, opts6):
        labels = [l for _, l in opts]
        check(all(any(l.startswith(c) for c in "①②③④⑤⑥⑦⑧⑨⑩") for l in labels)
              and len(set(l[0] for l in labels)) == len(labels),
              "配置方式编号连续不重复：%s" % labels)

    print("== 6c. 配置方式标题与「当前选择」显示 ==")
    _int_opts = typecatalog.fill_options(m2, "f", "int")
    _html = ui.options_html(_int_opts, "tag:ref")
    check("font-size:15px" in _html, "配置方式标题字号加大")
    check("当前选择" in _html and "#15803D" in _html, "当前选择标绿")
    check(_html.count("当前选择") == 1, "只有当前选择那一项被标出")
    check("当前选择" not in ui.options_html(_int_opts, "不存在的方式"),
          "选中的方式不在列表里时不乱标")

    print("== 6b. 标签类配置方式与右侧说明一一对应 ==")
    for t in ("int", "Quality", "list,Reward", "map,int,Reward"):
        style_keys = [k for k, _ in typecatalog.fill_style_options(m2, t)]
        opt_keys = [o["key"] for o in typecatalog.fill_options(m2, "f", t)]
        check(set(style_keys) <= set(opt_keys),
              "%s：下拉里每一项都在「配置方式」说明里有对应：%s ⊂ %s"
              % (t, style_keys, opt_keys))
    check("flags" not in [k for k, _ in typecatalog.fill_style_options(m2, "Quality")],
          "非 flags 枚举不提供「标志位组合」")
    _flags_enum = {o["key"]: o for o in typecatalog.fill_options(m2, "f", "Quality")}
    check("仅当" in _flags_enum["flags"]["title"],
          "非 flags 枚举的标志位示例标明前提：%s" % _flags_enum["flags"]["title"])
    fo = {o["key"]: o for o in typecatalog.fill_options(m2, "f", "int")}
    check(fo["tag:ref"]["grid"][1][2] == "int#ref=TbItem",
          "引用校验示例类型串：%s" % fo["tag:ref"]["grid"][1][2])
    fs = {o["key"]: o for o in typecatalog.fill_options(m2, "f", "list,Reward")}
    check("#size=4" in fs["tag:size"]["grid"][1][2],
          "长度约束示例类型串：%s" % fs["tag:size"]["grid"][1][2])

    print("== 7. 容器示例完整性（四种读法）==")
    fl = typecatalog.fill_options(m2, "items", "list,Reward")
    titles = " ".join(o["title"] for o in fl)
    for kw in ("流式", "sep", "列限定", "多行"):
        check(kw in titles, f"容器示例包含「{kw}」")
    for o in fl:
        g = o["grid"]
        check(len(g) >= 4 and str(g[0][0]).startswith("##"),
              f"「{o['title'][:14]}」表头完整（{len(g)} 行）")
        ncol = len(g[0])
        check(all(len(r) == ncol for r in g), f"「{o['title'][:14]}」各行列数一致")
        check(any(str(r[0]) == "##type" for r in g), f"「{o['title'][:14]}」含 ##type 行")

    print("== 8. 模板按配置方式生成 ==")
    t = TableDef(full_name="TbReward", value_type="Reward",
                 read_schema_from_file="false", input="reward.xlsx",
                 index="id", mode="map", group="c,s", comment="奖励表")
    m2.tables.append(t)
    fields = [FieldDef("id", "int", "", "ID"),
              FieldDef("items", "list,Reward", "", "奖励内容"),
              FieldDef("big", "Reward", "", "大奖")]
    p = template.default_input_path(m2, t)
    for style in ("sep", "stream", "column", "multiline"):
        for f in fields:
            f.fill_style = "auto"
        fields[1].fill_style = style
        fields[2].fill_style = style if style != "multiline" else "sep"
        ok, msg = template.generate_data_xlsx(
            m2, t, fields, p.replace(".xlsx", f"_{style}.xlsx"), style="auto")
        if not ok:
            check(False, f"模板生成失败({style})：{msg}")
            continue
        ws = load_workbook(p.replace(".xlsx", f"_{style}.xlsx")).worksheets[0]
        check(str(ws.cell(row=1, column=1).value).startswith("##"),
              f"{style}：A1 以 ## 开头")
    fields[1].fill_style, fields[2].fill_style = "auto", "auto"

    print("== 9. 静态自检 ==")
    issues = runner.validate_project(m2, {"TbReward": fields})
    check(not [i for i in issues if i["level"] == "error"], "自检无 error：%s" % issues)
    bad = [FieldDef("x", "list,int?", "", "")]
    check(any(i["level"] == "error" for i in typecatalog.validate_type(m2, "list,int?")),
          "拦住容器元素可空")

    print("== 10. 知识点 ==")
    check(len(knowledge.KNOWLEDGE) >= 20, "知识点 %d 条" % len(knowledge.KNOWLEDGE))
    check(bool(knowledge.search("多态")) and bool(knowledge.search("sep")), "检索可用")

    print("== 11. 界面构造（离屏）==")
    os.environ["QT_QPA_PLATFORM"] = "offscreen"
    from PySide6.QtWidgets import QApplication
    from PySide6.QtCore import QSettings
    app = QApplication.instance() or QApplication(sys.argv)
    # 界面测试里别让弹窗卡住脚本：一律自动回答
    from PySide6.QtWidgets import QMessageBox as _MB
    _MB.information = staticmethod(lambda *a, **k: _MB.Ok)
    _MB.warning = staticmethod(lambda *a, **k: _MB.Ok)
    _MB.critical = staticmethod(lambda *a, **k: _MB.Ok)
    _MB.question = staticmethod(lambda *a, **k: _MB.Yes)
    QSettings("LubanEditor", "LubanEditor").clear()   # 保证从「无记忆」状态开始
    app.setStyleSheet(ui.QSS)
    w = ui.MainWindow()
    check(w.stack.currentWidget() is w.home, "未指定目录时显示首页")
    check(all(getattr(w.home, n, None) is not None
              for n in ("bt_open", "bt_create", "bt_init")),
          "首页有「打开 / 创建 / 初始化」按钮对象")
    check(w.home.bt_open.text() == "打开工程" and w.home.bt_create.text() == "创建工程",
          "首页按钮名去掉「Luban」")
    check(w.home.bt_init.isHidden(), "首页隐藏「初始化工程」按钮")
    check(hasattr(w.home, "lst_recent"), "首页有「最近打开」列表")
    check(not w.home.lst_recent.isVisibleTo(w.home), "没有最近记录时列表隐藏")

    print("== 12. 打开工程：窗口标题 / 左侧目录树 / 只读预览 ==")
    w.open_project(ROOT)
    check(w.stack.currentWidget() is w.work, "打开工程后切到工作区")
    check(not w.toolbar.isHidden() and w.toolbar.actions(), "工具条显示且有按钮")
    check(w.windowTitle() == "LubanEditor",
          "窗口标题固定，只显示工具名：%s" % w.windowTitle())
    check(w.bt_folder.text() == "打开", "工程目录区右侧有「打开」按钮")
    # 界面精简：四个入口不再露出来，代码还在（新建表后自动生成数据表模板要用）
    _vis = [a.text() for a in w.toolbar.actions() if a.isVisible()]
    check(all(t not in _vis for t in ui.HIDDEN_TOOLBAR_ACTIONS),
          "隐藏「导出数据表模板 / 静态自检 / 生成并校验 / 设置」：%s" % _vis)
    check("知识点" not in _vis and "类型助手" not in _vis,
          "类型助手 / 知识点移入「帮助」菜单：%s" % _vis)
    check(hasattr(w, "help_btn") and w.help_btn.text() == "帮助",
          "工具条上有「帮助」按钮")
    check(all(hasattr(w, attr) for attr in ("act_export", "act_check", "act_gen",
                                            "act_settings")),
          "隐藏的入口其实现还在（内部流程用）")
    check(w.tree.topLevelItemCount() >= 1
          and w.tree.topLevelItem(0).text(0) == os.path.basename(ROOT),
          "左侧树根 = 工程目录名：%s" % w.tree.topLevelItem(0).text(0))
    dnode = w._find_dir_item(w.model.data_dir)
    check(dnode is not None, "目录树里能找到 Data 目录节点")
    snode = w._find_dir_item(os.path.join(ROOT, "Schemas"))
    check(snode is not None, "目录树里能找到 Schemas 目录节点")
    snames = [snode.child(i).text(0) for i in range(snode.childCount())]
    check(any(n.startswith("__tables__.xlsx") for n in snames),
          "Schemas 节点下能看到 __tables__.xlsx：%s" % snames)
    check(any("[定义]" in n for n in snames), "定义文件带 [定义] 标记")

    check(w.editor_stack.currentWidget() is w.page_file, "打开工程后默认是概览（只读页）")
    check("工程概览" in w.page_file.toPlainText(), "概览页有内容")

    # 定义文件 -> 只读摘要
    w._open_file(w.model.table_file)
    check(w.editor_stack.currentWidget() is w.page_file, "点定义文件走只读页")
    check("只读" in w.page_file.toPlainText(), "定义文件页标注只读")

    # 非 Data 目录的普通文件 -> 只读预览
    note = os.path.join(ROOT, "note_check.txt")
    with open(note, "w", encoding="utf-8") as f:
        f.write("hello luban\n第二行")
    w._open_file(note)
    check(w.editor_stack.currentWidget() is w.page_file, "普通文件走只读页")
    check("hello luban" in w.page_file.toPlainText(), "只读页显示文件内容")

    w.open_project(empty)
    check(w.editor_stack.currentWidget() is w.page_init,
          "未初始化目录：预览区中间显示「初始化工程」")
    check(hasattr(w, "bt_init_center") and w.bt_init_center.text() == "初始化工程",
          "预览区中间有「初始化工程」按钮")
    w.open_project(ROOT)
    check(w.editor_stack.currentWidget() is not w.page_init,
          "已初始化目录不显示初始化页")

    print("== 12b. 数据文件 -> 配置表 的匹配（回归：点开必须可编辑）==")
    # (1) input 留空：Luban 取表名。这类表以前匹配不到自己的数据文件 -> 点开是只读。
    t_blank = TableDef(full_name="TbBlank", value_type="Blank",
                       read_schema_from_file="true", input="",
                       index="id", mode="map", comment="input 留空的表")
    w.model.tables.append(t_blank)
    w.fields["TbBlank"] = [FieldDef("id", "int", "", "ID")]
    blank_path = os.path.join(w.model.data_dir, "TbBlank.xlsx")
    wb = Workbook()
    wb.active.append(["##var", "id"])
    wb.save(blank_path)
    w._reindex_inputs()
    w._refresh_tree()
    check(w._table_for_file(blank_path) == "TbBlank", "input 留空也能按表名匹配到数据文件")
    check(w._file_tag(blank_path) == "input 留空的表", "配置文件显示注释而不是 [配置表]")
    w._open_file(blank_path)
    check(w.editor_stack.currentWidget() is w.page_preview,
          "点配置表文件先进入只读预览")
    check(w._pending_table == "TbBlank", "预览记下了表名：%s" % w._pending_table)
    w._on_edit_table()
    check(w.current == "TbBlank", "编辑器载入的是那张表：%s" % w.current)
    dnode = w._find_dir_item(w.model.data_dir)
    same = [dnode.child(i).text(0) for i in range(dnode.childCount())
            if dnode.child(i).text(0).startswith("TbBlank.xlsx")]
    check(len(same) == 1, "同一个数据文件不再重复出现「[数据表未生成]」：%s" % same)

    # (2) input 写了别的相对路径 / 大小写不同 -> 按文件名兜底
    t_odd = TableDef(full_name="TbOdd", value_type="Odd",
                     read_schema_from_file="true", input="sub_dir/ODD.XLSX",
                     index="id", mode="map")
    r_before = list(w.model.tables)
    w.model.tables = [t_odd]
    odd_path = os.path.join(w.model.data_dir, "odd.xlsx")
    wb = Workbook()
    wb.active.append(["##var", "id"])
    wb.save(odd_path)
    w._reindex_inputs()
    check(w._table_for_file(odd_path) == "TbOdd", "input 路径/大小写不一致时按文件名兜底")
    w.model.tables = r_before
    w._reindex_inputs()
    w._refresh_tree()

    # (3) Data 里没被声明成配置表的 xlsx -> 只读，但要说清楚原因与出路
    stray = os.path.join(w.model.data_dir, "not_declared.xlsx")
    wb = Workbook()
    wb.active.append(["随便什么", "内容"])
    wb.save(stray)
    w._open_file(stray)
    check(w.editor_stack.currentWidget() is w.page_file, "未声明的 xlsx 走只读页")
    txt = w.page_file.toPlainText()
    check("没有在表定义文件里声明成配置表" in txt, "只读页说明「为什么不能编辑」")
    check("__tables__.xlsx" in txt and "tables.xml" in txt,
          "只读页列出当前工程的表定义文件（xlsx + xml 都列）：%s" % txt[:120])
    check("新建表" in txt, "只读页给出「怎么变成可编辑」的下一步")
    os.remove(stray)

    # (4) 精确匹配优先；只靠同名（且撞名）时不猜，宁可只读
    t_a = TableDef(full_name="TbA", value_type="A", input="same.xlsx",
                   index="id", mode="map")
    t_b = TableDef(full_name="TbB", value_type="B", input="dup/same.xlsx",
                   index="id", mode="map")
    r_keep = list(w.model.tables)
    w.model.tables = [t_a, t_b]
    w._reindex_inputs()
    check(w._table_for_file(os.path.join(w.model.data_dir, "same.xlsx")) == "TbA",
          "声明里的路径精确命中时按声明走")
    check(w._table_for_file(os.path.join(w.model.data_dir, "third", "same.xlsx")) == "",
          "只能靠同名匹配、且撞名时不乱猜")
    w.model.tables = r_keep
    w._reindex_inputs()
    w._refresh_tree()

    # (5) 表列表为空时，写盘必须拒绝 —— 否则整张表定义会被抹空，全工程瞬间变只读
    import tempfile
    tdir = tempfile.mkdtemp(prefix="luban_keep_")
    projectwizard.create_project(tdir, ["xlsx"])
    mprobe = ProjectModel()
    mprobe.root = tdir
    mprobe.table_file = os.path.join(tdir, "Schemas", "__tables__.xlsx")
    mprobe.tables = [TableDef(full_name="T1", value_type="A", input="t1.xlsx",
                              index="id", mode="map"),
                     TableDef(full_name="T2", value_type="B", input="t2.xlsx",
                              index="id", mode="map")]
    ok, msg = exceldef.save_tables(mprobe)
    check(ok, "先把两张表写进定义文件：%s" % msg)
    check(load_workbook(mprobe.table_file).worksheets[0].max_row == 5,
          "定义文件有 2 行表（3 行表头）")
    mprobe.tables = []
    ok2, msg2 = exceldef.save_tables(mprobe)
    check(not ok2 and "全部删掉" in msg2, "表列表为空时拒绝保存：%s" % msg2)
    check(load_workbook(mprobe.table_file).worksheets[0].max_row == 5,
          "被拒后定义文件原样保留（仍是 2 行表）")
    shutil.rmtree(tdir, ignore_errors=True)

    print("== 13. 配置表编辑：标红 / 保存 / 取消 ==")
    t2 = TableDef(full_name="TbDrop", value_type="Drop", read_schema_from_file="true",
                  input="drop.xlsx", index="id", mode="map", comment="掉落表")
    w.model.tables.append(t2)
    w.fields["TbDrop"] = [FieldDef("id", "int", "", "ID"),
                          FieldDef("reward", "Reward", "", "奖励"),
                          FieldDef("list2", "list,Reward", "", "奖励列表")]
    # 真造一个数据文件，走「点文件 -> 打开编辑器」这条真实链路
    drop_path = os.path.join(w.model.data_dir, "drop.xlsx")
    wb = Workbook()
    wb.active.append(["##var", "id", "reward", "list2"])
    wb.save(drop_path)
    w._reindex_inputs()
    w._refresh_tree()
    check(w._table_name_of_file(drop_path) == "TbDrop", "数据文件能映射回表名")
    check(w._file_tag(drop_path.strip()) == "掉落表", "文件树显示配置文件的注释而不是 [配置表]")

    w._open_file(drop_path)
    check(w.editor_stack.currentWidget() is w.page_preview, "点配置表文件先进入预览")
    w._on_edit_table()
    check(w.editor_stack.currentWidget() is w.page_table, "点「编辑」进入可编辑")
    check(w.tbl_fields.rowCount() == 3, "字段表载入 3 行")
    check(not w.bt_save_ed.isEnabled(), "刚载入时「保存」不可用（没有改动）")
    check(w.bt_save_ed.parent() is not None, "保存/取消在底部按钮条里")
    check(not w.bar_editor.isVisibleTo(w.page_table), "没有改动时整条保存栏都不显示")

    cb = w.tbl_fields.cellWidget(1, ui.COL_TYPE)
    check("Reward" in cb.currentText(), "类型框显示中文名：%s" % cb.currentText())
    cb2 = w.tbl_fields.cellWidget(2, ui.COL_STYLE)
    check(cb2.count() >= 3, "配置方式下拉有 %d 项" % cb2.count())
    check(w.tbl_fields.cellWidget(1, ui.COL_SUB).findChildren(ui.QLabel)[0].text() == "—",
          "非容器字段的「子类型」格显示 —")

    # 改「配置方式」：先进缓冲区并标红，不碰已保存数据
    cb2.setCurrentIndex(cb2.findData("multiline"))
    check(w._edit_fields[2].fill_style == "multiline", "改动进入缓冲区")
    check(w.fields["TbDrop"][2].fill_style != "multiline", "未保存前不动已保存数据")
    check(w.bt_save_ed.isEnabled() and w.lab_dirty.isVisibleTo(w.page_table),
          "出现未保存提示且「保存」可用")
    check(w.bt_save_ed.isVisibleTo(w.page_table) and w.bt_cancel_ed.isVisibleTo(w.page_table),
          "有改动时保存 / 取消按钮才出现")
    check(bool(w.tbl_fields.cellWidget(2, ui.COL_STYLE).styleSheet()),
          "改动过的「配置方式」控件标红")

    # 改单元格文字：字段名 / 注释 直接同步进缓冲区并标红
    w.tbl_fields.item(2, ui.COL_COMMENT).setText("改过的注释")
    check(w._edit_fields[2].comment == "改过的注释", "单元格改动进入缓冲区")
    red = w.tbl_fields.item(2, ui.COL_COMMENT).foreground().color().name()
    check(red == "#b42318", "改动过的单元格标红：%s" % red)

    # 改「类型」-> 可空，类型串被改写
    cb0 = w.tbl_fields.cellWidget(0, ui.COL_STYLE)
    cb0.setCurrentIndex(cb0.findData("nullable"))
    check(w._edit_fields[0].type == "int?",
          "选可空后类型变成 int?（实际 %s）" % w._edit_fields[0].type)

    # 表单改动也进缓冲区
    w.ed_input.setText("changed.xlsx")
    check(w._form["input"] == "changed.xlsx", "表定义表单改动进入缓冲区")
    check(w._base_form["input"] == "drop.xlsx", "缓冲不影响基线")

    # 「取消修改」丢弃全部缓冲
    w._cancel_editor()
    check(w._edit_fields[0].type == "int" and w._form["input"] == "drop.xlsx",
          "「取消修改」丢弃全部缓冲改动")
    check(w._edit_fields[2].comment == "奖励列表", "取消后字段注释回到原值")
    check(not w.bt_save_ed.isEnabled(), "取消后「保存」重新变灰")
    check(w.tbl_fields.item(2, ui.COL_COMMENT).foreground().color().name() != "#b42318",
          "取消后红色标记被清掉")

    # 再改一次并提交（只验证提交逻辑，不落盘以免弹窗）
    cb2 = w.tbl_fields.cellWidget(2, ui.COL_STYLE)
    cb2.setCurrentIndex(cb2.findData("multiline"))
    check(w._commit_editor(), "提交缓冲区")
    check(w.fields["TbDrop"][2].fill_style == "multiline", "提交后已保存数据被更新")
    w._repaint_dirty()
    check(not w.bt_save_ed.isEnabled(), "提交后基线刷新，「保存」变灰")

    print("== 13b. 容器类型与子类型下拉 ==")
    # 类型下拉里必须有容器（以前只能手写 list,Reward）
    lv_row = 0                    # TbDrop 第 0 行是 id（int）
    tc0 = w.tbl_fields.cellWidget(lv_row, ui.COL_TYPE)
    tc_names = [tc0.itemText(i) for i in range(tc0.count())]
    for stub in ("list,int（列表）", "array,int（数组）", "set,int（集合）",
                 "map,int,string（映射）"):
        check(stub in tc_names, "类型下拉含容器：%s" % stub)
    check(tc_names[-1] == ui.ADD_NEW_TYPE, "「新建自定义类型」仍在下拉最后")

    # 选 list -> 类型串带默认元素；子类型格出现下拉；配置方式换成容器那一套
    tc0.setCurrentIndex(tc_names.index("list,int（列表）"))
    check(w._edit_fields[lv_row].type == "list,int",
          "选容器后类型串变成 list,int（%s）" % w._edit_fields[lv_row].type)
    subs = w.tbl_fields.cellWidget(lv_row, ui.COL_SUB).findChildren(ui.ValueCombo)
    check(len(subs) == 1, "list 给一个「元素」子类型下拉：%d" % len(subs))
    check([l.text() for l in w.tbl_fields.cellWidget(lv_row, ui.COL_SUB)
           .findChildren(ui.QLabel)] == ["元素"], "list 的子类型标「元素」")
    check(subs[0].code() == "int", "子类型默认与类型串一致：%s" % subs[0].code())
    st0 = w.tbl_fields.cellWidget(lv_row, ui.COL_STYLE)
    check(st0.findData("sep") >= 0 and st0.findData("multiline") >= 0,
          "换成容器后「配置方式」是容器那一套（%d 项）" % st0.count())
    check(bool(tc0.styleSheet()), "改过类型的行标红")

    # 改元素类型 -> 类型串跟着改，并回显到类型框
    subs[0].set_code("Reward")
    check(w._edit_fields[lv_row].type == "list,Reward",
          "换元素后类型串变成 list,Reward（%s）" % w._edit_fields[lv_row].type)
    check(tc0.currentText() == "list,Reward（列表）",
          "类型框回显新类型：%s" % tc0.currentText())

    # map 需要 key + value 两个子类型
    tc0.setCurrentIndex([i for i, n in enumerate(tc_names)
                         if n.startswith("map,int,string")][0])
    check(w._edit_fields[lv_row].type == "map,int,string",
          "选 map 后类型串 = %s" % w._edit_fields[lv_row].type)
    subs = w.tbl_fields.cellWidget(lv_row, ui.COL_SUB).findChildren(ui.ValueCombo)
    check(len(subs) == 2, "map 给 key + value 两个子类型下拉：%d" % len(subs))
    check([l.text() for l in w.tbl_fields.cellWidget(lv_row, ui.COL_SUB)
           .findChildren(ui.QLabel)] == ["键", "值"], "map 的子类型标「键 / 值」")
    subs[0].set_code("string")
    subs[1].set_code("Reward")
    check(w._edit_fields[lv_row].type == "map,string,Reward",
          "改 key / value 后类型串 = %s" % w._edit_fields[lv_row].type)
    check(any(t.startswith("Reward") for c in subs for t in
              [c.itemText(i) for i in range(c.count())]),
          "子类型下拉与「类型」下拉同源（含工程里的结构）")

    # 带上 #size 这类容器标签时，改子类型不能把标签弄丢
    tc0.setCurrentText("map,int,Reward（映射）")
    tc0.lineEdit().editingFinished.emit()
    check(w._edit_fields[lv_row].type == "map,int,Reward",
          "手写容器类型生效：%s" % w._edit_fields[lv_row].type)
    st0 = w.tbl_fields.cellWidget(lv_row, ui.COL_STYLE)
    st0.setCurrentIndex(st0.findData("tag:size"))
    tagged = w._edit_fields[lv_row].type
    check("#size=4" in tagged, "容器可以挂长度约束：%s" % tagged)
    subs = w.tbl_fields.cellWidget(lv_row, ui.COL_SUB).findChildren(ui.ValueCombo)
    subs[0].set_code("string")
    check("#size=4" in w._edit_fields[lv_row].type
          and w._edit_fields[lv_row].type.endswith("string,Reward"),
          "改子类型时保留容器上的 #size：%s" % w._edit_fields[lv_row].type)

    # 带标签的元素类型（(Reward#sep=,) 表示「结构在一个格子里用 , 分隔子字段」）
    # 在下拉里必须显示成可读名，而不是把原始写法丢给用户看
    lab = typecatalog.elem_label(w.model, "(Reward#sep=,)")
    check(lab.startswith("Reward") and "#" not in lab and "分隔" in lab,
          "带标签的元素类型显示成可读名：%s" % lab)
    cbx = ui.ValueCombo([("int", "int（整数）")])
    cbx.set_code("(Reward#sep=,)", typecatalog.elem_label(w.model, "(Reward#sep=,)"))
    check(cbx.code() == "(Reward#sep=,)" and "#" not in cbx.currentText(),
          "子类型框显示可读名、取值仍是完整写法：%s / %s"
          % (cbx.currentText(), cbx.code()))

    # 点到另一行字段，右侧「配置方式」要跟着换行（否则选 A 行看 B 行的示例）
    w.tbl_fields.selectRow(1)
    check("reward" in w.view_opts.toPlainText()
          and "字段 reward" in w.view_opts.toPlainText().replace("\u00a0", " "),
          "点选另一行后右侧换成该字段的配置方式：%s"
          % w.view_opts.toPlainText().splitlines()[:3])

    print("== 13c. 分组下拉（带中文名）==")
    gcell = w.tbl_fields.cellWidget(0, ui.COL_GROUP)
    gcb = gcell.findChild(ui.ValueCombo) if gcell is not None else None
    check(gcb is not None, "字段「分组」格是下拉框，不再是纯文本框")
    gopts = [gcb.itemText(i) for i in range(gcb.count())] if gcb else []
    check("c（客户端）" in gopts and "s（服务器）" in gopts,
          "分组下拉带中文名：%s" % gopts[:4])
    check(gcb is not None and gcb.isEditable(), "分组下拉可手写（多个分组 c,s）")
    check(w.cb_group.count() >= 4
          and "c（客户端）" in [w.cb_group.itemText(i) for i in range(w.cb_group.count())],
          "表定义的「导出分组」下拉也带中文名")
    before = w._edit_fields[0].group
    gcb.set_code("c")
    check(w._edit_fields[0].group == "c", "选分组后写回字段（%s → c）" % (before or "空"))
    check("#B42318" in w.tbl_fields.cellWidget(0, ui.COL_GROUP).styleSheet(),
          "改了分组会标红待保存")

    print("== 13d. 单元格格式（stream / lite / json / lua）==")
    check(not ui.MainWindow._is_complex(w.model, "int"),
          "基础类型不需要单元格格式")
    check(ui.MainWindow._is_complex(w.model, "list,int")
          and ui.MainWindow._is_complex(w.model, "Reward"),
          "容器 / 结构 才有单元格格式")
    fcell = w.tbl_fields.cellWidget(lv_row, ui.COL_FMT)
    fcb = fcell.findChild(ui.ValueCombo) if fcell is not None else None
    check(fcb is not None, "容器行的「单元格格式」是下拉")
    fopts = [(fcb.itemText(i), fcb.itemData(i)) for i in range(fcb.count())] if fcb else []
    check([c for _t, c in fopts] == ["", "lite", "json", "lua"],
          "四个选项：stream（默认）/ lite / json / lua -> %s" % [t for t, _c in fopts])
    fcb.set_code("lite")
    check(w._edit_fields[lv_row].fmt == "lite",
          "选 lite 写回字段：%s" % w._edit_fields[lv_row].fmt)
    check(not w.tbl_fields.cellWidget(lv_row, ui.COL_STYLE).isEnabled(),
          "选了紧凑格式后「配置方式」自动禁用（两边不打架）")
    cols_fmt = template.build_cols(w.model, [w._edit_fields[lv_row]], default="sep")
    check(len(cols_fmt) == 1 and cols_fmt[0].top.endswith("#format=lite"),
          "#format= 写在字段名上：%s" % (cols_fmt[0].top if cols_fmt else ""))
    check(cols_fmt and cols_fmt[0].sample.startswith("{{"),
          "lite 示例是花括号：%s" % (cols_fmt[0].sample if cols_fmt else ""))
    fcb.set_code("json")
    cols_json = template.build_cols(w.model, [w._edit_fields[lv_row]], default="sep")
    check(cols_json and cols_json[0].top.endswith("#format=json")
          and cols_json[0].sample.startswith('[["'),
          "map 的 json 示例按 [key, value] 成对：%s"
          % (cols_json[0].sample if cols_json else ""))
    fcb.set_code("")
    check(w._edit_fields[lv_row].fmt == "", "选回 stream 就是默认（不写 #format）")
    check(w.tbl_fields.cellWidget(lv_row, ui.COL_STYLE).isEnabled(),
          "回到 stream 后「配置方式」恢复可用")

    w._cancel_editor()

    print("== 14. 对话框 ==")
    # 类型助手：四区布局 + 常用标签随类型变化 + 通用 bean / 多态bean
    dlg = ui.TypePickerDialog(w, w.model, "reward", "int", "TbDrop")
    check(dlg.value() == "int", "类型助手初值正确")
    check(dlg.gb_tag.title().startswith("常用标签"), "右上区是「常用标签」")
    check(dlg.lst.count() > 0, "左侧类型列表非空")
    texts = [dlg.lst.item(i).text().strip() for i in range(dlg.lst.count())]
    check("bean（通用结构）" in texts, "类型列表有通用 bean")
    check("多态bean（多态结构）" in texts, "类型列表有多态bean")
    check(not any(t.startswith("＋") for t in texts), "类型列表去掉了「新建自定义类型」")
    check("ref" in dlg._vis_keys, "基础类型的常用标签含引用校验")
    dlg.ck_null.setChecked(True)
    dlg._rebuild()
    check(dlg.value() == "int?", "勾可空后变 int?（%s）" % dlg.value())
    dlg.ck_null.setChecked(False)
    dlg._rebuild()
    check(dlg.value() == "int", "取消勾选后恢复（%s）" % dlg.value())
    for _ in range(3):
        dlg.tag_cks["ref"].setChecked(True)
        dlg._rebuild()
    check(dlg.value().count("#ref=") == 1, "反复勾选不会重复追加：%s" % dlg.value())

    # 换成容器：标签集合变化（ref 没了、出现 sep/size），可空也不给
    dlg.ed.setText("list,int")
    check("sep" in dlg._vis_keys and "size" in dlg._vis_keys,
          "容器的常用标签含 sep / size：%s" % dlg._vis_keys)
    check("ref" not in dlg._vis_keys, "容器不再提供引用校验标签")
    check(not dlg._null_vis, "容器不给「可空」")
    check("<h4>list,int</h4>" in dlg.view.toHtml() or "list,int" in dlg.view.toHtml(),
          "右下角展示该类型的配置示例")

    # 通用 bean / 多态bean：纯讲解，不报「未知类型」
    dlg.ed.setText("bean")
    dlg.lst.setCurrentItem(dlg.lst.item(texts.index("bean（通用结构）")))
    dlg._refresh()
    check("通用结构" in dlg.view.toHtml(), "选通用 bean 显示通用结构讲解")
    check("color:#c00" not in dlg.view.toHtml(), "通用 bean 不再给出「未知类型」报错")
    dlg.lst.setCurrentItem(dlg.lst.item(texts.index("多态bean（多态结构）")))
    dlg._refresh()
    check("多态" in dlg.view.toHtml(), "选多态bean 显示多态讲解")

    # 手写类型串也能正确判定
    dlg.ed.setText("Reward")
    check(dlg._eff_kind()[0] == "bean", "手写结构名判定为 bean")
    dlg.ed.setText("list,Reward")
    check(dlg._eff_kind()[0] == "container", "手写容器判定为容器")

    # 建工程对话框
    d2 = ui.CreateProjectDialog(w, ROOT)
    check(d2.result()[1] == ["xlsx"], "建工程默认勾选 xlsx")
    check("创建" in d2.windowTitle(), "创建模式标题：%s" % d2.windowTitle())
    d2b = ui.CreateProjectDialog(w, ROOT, "init")
    check("初始化" in d2b.windowTitle(), "初始化模式标题：%s" % d2b.windowTitle())
    d3 = ui.CreateTypeDialog(w, w.model, "TbDrop")
    check(d3.rb_bean.isChecked() and d3.rb_common.isChecked(),
          "新建类型：默认同时选中「结构」和「公共」")
    d3.rb_table.setChecked(True)
    check(d3.rb_bean.isChecked(),
          "点「归属」不会把上面的「结构」取消掉")
    check(d3.rb_table.isChecked() and not d3.rb_common.isChecked(),
          "归属组内部互斥：选了当前表专用就不是公共")
    d3.rb_enum.setChecked(True)
    check(d3.rb_enum.isChecked() and not d3.rb_bean.isChecked(),
          "结构 / 枚举 组内部互斥")
    check(d3.rb_table.isChecked(), "切到枚举不影响已选的「归属」")
    check(len(d3.grp_kind.buttons()) == 2 and len(d3.grp_scope.buttons()) == 2,
          "两组单选框各自成组：%d / %d"
          % (len(d3.grp_kind.buttons()), len(d3.grp_scope.buttons())))
    check(d3.bt_ok.text() == "确定" and d3.bt_cancel.text() == "取消",
          "底部按钮是「确定 / 取消」：%s / %s" % (d3.bt_ok.text(), d3.bt_cancel.text()))
    check(d3.bt_ok.objectName() == "footprimary"
          and d3.bt_cancel.objectName() == "foot",
          "底部按钮用加大的居中样式：%s / %s"
          % (d3.bt_ok.objectName(), d3.bt_cancel.objectName()))
    d3.rb_bean.setChecked(True)
    d3.rb_common.setChecked(True)
    d3.ed_name.setText("MyBean")
    check(d3.full_name() == "MyBean", "归属公共时类型名不加前缀")
    d3.rb_table.setChecked(True)
    check(d3.full_name().endswith(".MyBean"), "归属当前表时加模块前缀：%s" % d3.full_name())

    print("== 14a. 发布工程面板 ==")
    acts = [a.text() for a in w.toolbar.actions()]
    check("发布工程" in acts, "工具条有「发布工程」入口：%s" % acts)
    check(acts.index("发布工程") == acts.index("打开工程") + 1,
          "「发布工程」紧跟在「打开工程」后面：%s" % acts)
    dp = ui.PublishDialog(w, w.model, w.settings)
    dp.show()
    # 参数分基础 / 高级：基础默认可见，高级默认折叠，点「更多」才展开
    check(dp.gb_a.isHidden() and dp.bt_more.text().startswith("▸"),
          "高级参数默认折叠")
    dp.bt_more.click()
    check(dp.gb_a.isVisible() and dp.bt_more.text().startswith("▾"),
          "点「更多」才展开高级参数")
    for k in ("exeMode", "exe", "conf", "target", "codeTarget", "dataTarget",
              "outputCodeDir", "outputDataDir"):
        check(k in dp._fields, "有基础参数：%s" % k)
    for k in ("schemaCollector", "pipeline", "includeTag", "excludeTag",
              "outputTable", "variant", "timeZone", "outputSaver",
              "cleanUpOutputDir", "fileExt", "jsonCompact", "lineEnding",
              "l10nProvider", "l10nTextFile"):
        check(k in dp._fields, "有高级参数：%s" % k)
    for kk, wg in dp._fields.items():
        check(bool((wg.toolTip() or "").strip()), "参数「%s」有备注含义" % kk)
    check(len([p for p in publish.PARAMS if p.section == "basic"]) == 8,
          "基础参数 8 个（执行方式 + 官方模板脚本里的那几个）")
    check(isinstance(dp._fields["exeMode"], ui.ExeModeSwitch),
          "「执行方式」是 dll / exe 二选一")
    check(len(dp._fields["exeMode"].grp.buttons()) == 2, "只有 dll / exe 两个选项")
    check(len([p for p in publish.PARAMS if p.section == "advanced"]) >= 10,
          "高级参数有 %d 个" % len([p for p in publish.PARAMS
                                   if p.section == "advanced"]))
    cb_t = dp._fields["target"]
    check(all(t in [cb_t.itemText(i) for i in range(cb_t.count())]
              for t in m2.targets), "-t 下拉里是 luban.conf 的导出目标")
    for k in ("target", "codeTarget", "dataTarget"):
        check(isinstance(dp._fields[k], ui.QComboBox), "%s 用下拉设置" % k)
    dp._fields["dataTarget"].setEditText("lua")
    dp._fields["codeTarget"].setEditText("cs-bin")
    cmd = dp.ed_cmd.text()
    check("-d lua" in cmd and "-c cs-bin" in cmd,
          "改参数后命令预览同步：%s" % cmd)
    dp._copy_cmd()
    check(QApplication.clipboard().text() == dp.ed_cmd.text(),
          "「复制命令」把命令写进剪贴板")
    # 文件 / 目录类参数：有输入框 + 「选择…」按钮
    for k in ("exe", "conf", "outputCodeDir", "outputDataDir"):
        wg = dp._fields[k]
        box = wg.parent()
        check(box.findChild(ui.QPushButton) is not None,
              "「%s」支持选择：%s" % (k, "…"))

    print("== 14a2. 管理自定义类型：入口只在打开 Luban 工程时出现 ==")
    check("管理自定义类型" in acts, "工具条有「管理自定义类型」入口：%s" % acts)
    a_mt = [a for a in w.toolbar.actions() if a.text() == "管理自定义类型"][0]
    check(w.toolbar.widgetForAction(a_mt).objectName() != "deep",
          "「管理自定义类型」用普通色按钮")
    check(a_mt.isVisible(), "打开了 Luban 工程时入口可见")
    # 换一个「不是 Luban 工程」的目录：入口要自己藏起来
    ndir = os.path.join(tempfile.gettempdir(), "lubankit_noproj")
    shutil.rmtree(ndir, ignore_errors=True)
    os.makedirs(ndir, exist_ok=True)
    wn = ui.MainWindow()
    wn.open_project(ndir)
    check(not wn.action_manage_types.isVisible(),
          "打开的目录不是 Luban 工程时，入口隐藏")
    wn.close()

    print("== 14b. 分列示例：跨列字段要画成合并格 ==")
    # 复合字段横跨多列，续列不该看起来像独立空列
    fl_html = ui.options_html(typecatalog.fill_options(m2, "items", "list,Reward"), "stream")
    check("colspan=" in fl_html, "多列字段渲染成合并单元格")
    check("title='这一格跨" in fl_html, "合并格写明跨了几列")
    check("<b>黄色合并的格子</b>" in fl_html, "给出「合并格 = 一个字段占多列」的说明")
    check(ui.MULTI_SPAN_COLORS[0] in fl_html, "多列格子填成黄色（%s）" % ui.MULTI_SPAN_COLORS[0])
    _stream = {o["key"]: o for o in typecatalog.fill_options(m2, "items", "list,Reward")}
    g = _stream["stream"]["grid"]
    check(len(g[0]) == len(g[1]) == len(g[2]),
          "示例各行（var/type/group）列数一致：%s" % [len(r) for r in g])
    check(g[0][3] is typecatalog.PAD and g[2][3] is typecatalog.PAD,
          "容器流式示例：字段名与类型后面的续列标成 PAD（合并）")
    check(sum(1 for c in g[0] if c is not typecatalog.PAD) < len(g[0]),
          "容器流式示例确实占用了多列")
    _map_g = {o["key"]: o for o in typecatalog.fill_options(m2, "f", "map,int,Reward")}
    mg = _map_g["column"]["grid"]
    check(len(mg[0]) == 4 and mg[1][2] == "ch-zn" and mg[1][3] == "en-us",
          "map 多级表头示例不再被截断（key 列都在）：%s" % mg[1])
    check(all(len(r) == len(mg[0]) for r in mg), "map 示例各行列数一致")
    # 模板预览也要合并
    cols = template.build_cols(
        m2, [FieldDef("items", "list,Reward", "", "奖励", fill_style="stream")],
        default="auto")
    hg = template.header_grid(cols)
    check(len(hg[0]) == len(cols) + 1, "表头预览列数 = 实际列数：%d" % len(hg[0]))
    check(any(c is typecatalog.PAD for c in hg[0]), "表头预览里续列标成 PAD")
    tp = ui._html_table(hg)
    check("colspan=" in tp, "表头预览渲染出合并格")
    check("colspan=" not in ui._html_table([["##var", "id", "lv"], ["##type", "int", "int"]]),
          "单列字段不产生多余的合并格")

    print("== 15. 四区布局 / 新建表内联 ==")
    # 四区：顶部功能区 + 左(工程区) / 中(文件区) / 右(上下文区)
    check(not w.toolbar.isHidden(), "顶部功能区：工具条显示")
    acts = [a.text() for a in w.toolbar.actions()]
    check("类型助手" not in acts, "类型助手移入「帮助」菜单：%s" % acts)
    check("保存定义" not in acts, "顶部去掉了「保存定义」按钮")
    check(w.workspace.count() == 3, "工作区分三列：%d" % w.workspace.count())
    check(w.workspace.widget(0) is w.left_panel, "第 1 列 = 左侧工程区（标题 + 目录树）")
    check(w.workspace.widget(1) is w.editor_stack, "第 2 列 = 中间文件区")
    check(w.workspace.widget(2) is w.right_panel, "第 3 列 = 右侧上下文区")
    check([w.tabs.tabText(i) for i in range(w.tabs.count())]
          == ["配置方式", "模板预览"],
          "右侧两个页签：%s" % [w.tabs.tabText(i) for i in range(w.tabs.count())])

    # 中间区六页：编辑器 / 只读预览 / 新建表 / 配置表预览 / 初始化 / 自定义类型预览
    check(w.editor_stack.count() == 6, "中间区六页：%d" % w.editor_stack.count())
    check(w.editor_stack.widget(0) is w.page_table
          and w.editor_stack.widget(1) is w.page_file
          and w.editor_stack.widget(2) is w.page_new
          and w.editor_stack.widget(3) is w.page_preview
          and w.editor_stack.widget(4) is w.page_init
          and w.editor_stack.widget(5) is w.page_def,
          "中间区 = 编辑器 / 只读预览 / 新建表 / 配置表预览 / 初始化 / 自定义类型预览")
    check(not hasattr(ui, "NewTableDialog"), "新建表不再用独立窗口类")

    # 打开新建表：显示在中间区，不弹窗口
    w._open_new_table()
    check(w.editor_stack.currentWidget() is w.page_new, "新建表显示在中间区")
    check(w.nt_fields.rowCount() == 1, "新建表默认给一行字段")
    check(w.nt_vertical.isChecked() is False, "纵表模板默认不勾")
    check(w.bt_nt_create.text() == "创建", "新建表底部有「创建」按钮")
    # 表声明文件类型：默认 Excel Schema，切到 XML Schema 目标提示要跟着变
    check(w.rb_nt_xlsx.isChecked() and not w.rb_nt_xml.isChecked(),
          "表声明文件类型默认 Excel Schema")
    check(w._nt_def_type == "xlsx", "默认类型 = xlsx")
    check("__tables__.xlsx" in w.lab_nt_target.text(),
          "Excel Schema 提示写入 __tables__.xlsx：%s" % w.lab_nt_target.text())
    w.rb_nt_xml.setChecked(True)
    check(w._nt_def_type == "xml" and "tables.xml" in w.lab_nt_target.text(),
          "切到 XML Schema 后提示写入 tables.xml：%s" % w.lab_nt_target.text())
    w.rb_nt_xlsx.setChecked(True)
    check(w._nt_def_type == "xlsx", "能切回 Excel Schema")

    # 新建表的字段区要和「预览表」完全一致：7 列 + 容器/子类型 + 可新建类型 + 配置方式 + 助手
    check([w.nt_fields.horizontalHeaderItem(i).text() for i in range(w.nt_fields.columnCount())]
          == [w.tbl_fields.horizontalHeaderItem(i).text()
              for i in range(w.tbl_fields.columnCount())],
          "新建表字段列与预览表一致：%s"
          % [w.nt_fields.horizontalHeaderItem(i).text()
             for i in range(w.nt_fields.columnCount())])
    check([w.nt_fields.horizontalHeaderItem(i).text() for i in range(w.nt_fields.columnCount())]
          == ui.FIELD_COLS, "字段列 = 字段名/类型/子类型/配置方式/分组/注释/助手")
    nt_tc = w.nt_fields.cellWidget(0, ui.COL_TYPE)
    check(nt_tc.itemText(nt_tc.count() - 1) == ui.ADD_NEW_TYPE,
          "新建表的类型下拉含「新建自定义类型」")
    nt_names = [nt_tc.itemText(i) for i in range(nt_tc.count())]
    check(any(n.startswith("set,int") for n in nt_names), "新建表的类型下拉含容器")
    nt_sc = w.nt_fields.cellWidget(0, ui.COL_STYLE)
    check(nt_sc is not None and nt_sc.count() >= 6,
          "新建表有「配置方式」下拉且内容完整（%d 项）" % (nt_sc.count() if nt_sc else 0))
    check(isinstance(w.nt_fields.cellWidget(0, ui.COL_HELP), ui.QPushButton),
          "新建表有「助手」按钮")
    nt_btns = [b.text() for b in w.nt_fields.parent().findChildren(ui.QPushButton)
               if b.text().startswith(("+", "-", "↑", "↓"))]
    check(len(nt_btns) == 4, "新建表字段区有增删与上下移：%s" % nt_btns)

    # 新建表里改「配置方式」也要改写类型串，且右侧跟着变
    nt_sc.setCurrentIndex(nt_sc.findData("nullable"))
    check(w._nt_buf[0].type == "int?", "新建表选可空后类型串变成 int?（%s）" % w._nt_buf[0].type)
    nt_sc2 = w.nt_fields.cellWidget(0, ui.COL_STYLE)
    check(nt_sc2.findData("tag:ref") >= 0, "新建表也能选校验标签类配置方式")
    nt_sc2.setCurrentIndex(nt_sc2.findData("tag:ref"))
    check("#ref=" in w._nt_buf[0].type,
          "新建表选引用校验后类型串带 #ref（%s）" % w._nt_buf[0].type)
    w.nt_fields.selectRow(0)
    check("当前选择" in w.view_opts.toHtml(), "新建表右侧「配置方式」标出当前选择")
    check("新建表" in w.view_opts.toHtml(), "右侧说明是新表的字段")
    w._nt_add_field()
    check(w.nt_fields.rowCount() == 2 and w._nt_buf[1].name == "field_2",
          "新建表加字段同步缓冲区")
    # 新建表的子类型也能用：选容器 -> 元素下拉 -> 类型串跟着变
    nt_tc = w.nt_fields.cellWidget(1, ui.COL_TYPE)
    nt_names = [nt_tc.itemText(i) for i in range(nt_tc.count())]
    nt_tc.setCurrentIndex([i for i, n in enumerate(nt_names)
                           if n.startswith("array,int")][0])
    nt_subs = w.nt_fields.cellWidget(1, ui.COL_SUB).findChildren(ui.ValueCombo)
    check(len(nt_subs) == 1 and w._nt_buf[1].type == "array,int",
          "新建表选容器后带出子类型下拉（%s）" % w._nt_buf[1].type)
    nt_subs[0].set_code("Reward")
    check(w._nt_buf[1].type == "array,Reward",
          "新建表改元素类型后类型串 = %s" % w._nt_buf[1].type)
    # 类型下拉可以手写（容器 / 带标签的类型），回车或失焦时要落到字段上
    nt_le = w.nt_fields.cellWidget(1, ui.COL_TYPE)
    nt_le.setCurrentText("list,Reward")
    nt_le.lineEdit().editingFinished.emit()
    check(w._nt_buf[1].type == "list,Reward",
          "手写类型按下回车后进入缓冲区：%s" % w._nt_buf[1].type)
    check(w.nt_fields.cellWidget(1, ui.COL_STYLE).findData("sep") >= 0,
          "手写容器后「配置方式」跟着换成容器的选项")
    check(len(w.nt_fields.cellWidget(1, ui.COL_SUB).findChildren(ui.ValueCombo)) == 1,
          "手写容器后子类型下拉出现")
    nt_le.setCurrentText("int")
    nt_le.lineEdit().editingFinished.emit()
    check(w._nt_buf[1].type == "int", "改回 int 也能同步：%s" % w._nt_buf[1].type)
    check(len(w.nt_fields.cellWidget(1, ui.COL_SUB).findChildren(ui.ValueCombo)) == 0,
          "改回非容器后子类型下拉收起")
    w.nt_fields.selectRow(1)
    w._nt_move(-1)
    check([f.name for f in w._nt_buf] == ["field_2", "field_1"],
          "新建表可上下移：%s" % [f.name for f in w._nt_buf])
    w.nt_fields.selectRow(1)
    w._nt_del_field()
    check(w.nt_fields.rowCount() == 1 and len(w._nt_buf) == 1, "新建表可删字段")

    w.nt_input.setText("inline")
    w.nt_name.setText("TbInline")
    w.nt_mode.set_code("map")
    w.nt_index.setText("id")
    w.nt_fields.item(0, ui.COL_NAME).setText("id")
    w.nt_fields.cellWidget(0, ui.COL_TYPE).setCurrentText("int（整数）")
    tt = w._nt_result_table()
    check(tt.full_name == "TbInline" and tt.input == "inline.xlsx"
          and tt.index == "id" and tt.real_mode == "map",
          "内联表单算出 TableDef：%s / %s" % (tt.full_name, tt.input))
    ff = w._nt_result_fields()
    check(len(ff) == 1 and ff[0].name == "id" and ff[0].type == "int",
          "内联表单算出字段：%s" % [(f.name, f.type) for f in ff])

    # 取消 -> 回概览
    w._nt_cancel()
    check(w.editor_stack.currentWidget() is w.page_file, "取消新建后回到概览")

    # mode=one 自动清空主键
    w._open_new_table()
    w.nt_name.setText("TbSingle")
    w.nt_mode.set_code("one")
    w.nt_index.setText("id")
    check(w._nt_result_table().index == "", "mode=one 时自动清空主键")
    w._nt_cancel()

    # 保存只校验当前这张表：别的表有问题不该拦我
    real_vp, real_save = runner.validate_project, exceldef.save_tables
    real_show = w._show_issues
    shown = []
    try:
        runner.validate_project = lambda model, fields=None: [
            {"level": "warn", "table": "TbOther", "msg": "别的表的小提示"},
            {"level": "warn", "table": "TbDrop", "msg": "本表的小提示"},
        ]
        exceldef.save_tables = lambda model: (True, "ok")
        w._show_issues = lambda issues: shown.extend(issues)
        ok_save = w._save_definitions("TbDrop")
    finally:
        runner.validate_project, exceldef.save_tables = real_vp, real_save
        w._show_issues = real_show
    check(ok_save and shown and all(i["table"] == "TbDrop" for i in shown),
          "保存只校验/提示当前表：%s" % [i["table"] for i in shown])

    print("== 16. 删除配置表：声明与数据文件分开处理 ==")
    data_path = os.path.join(w.model.data_dir, "TbDrop.xlsx")
    if not os.path.exists(data_path):          # 合成工程里补一个占位文件
        with open(data_path, "wb") as f:
            f.write(b"")
    check(w.model.find_table("TbDrop") is not None, "删除前 TbDrop 在")
    check(os.path.exists(data_path), "删除前数据文件在：%s" % os.path.basename(data_path))
    w._delete_table("TbDrop", with_file=False)      # 跳过对话框，只删声明
    check(w.model.find_table("TbDrop") is None, "声明已从 model 里删除")
    check(os.path.exists(data_path), "「只删声明」保留数据文件")
    check("删除" in getattr(w, "_toast", None).text() if getattr(w, "_toast", None) else False,
          "删除后也有轻提示")

    print("== 17. 多列表头：合并单元格 / 黄色深浅 / 列序备注 ==")
    ddir = os.path.join(tempfile.gettempdir(), "lubankit_tpl")
    shutil.rmtree(ddir, ignore_errors=True)
    os.makedirs(ddir)
    tm = ProjectModel(data_dir=ddir)
    tm.beans.append(BeanDef(full_name="Vec3", fields=[
        FieldDef("x", "float"), FieldDef("y", "float"), FieldDef("z", "float")]))
    tm.beans.append(BeanDef(full_name="Item", fields=[
        FieldDef("id", "int"), FieldDef("num", "int")]))
    tf = [FieldDef("id", "int", "", "编号"),
          FieldDef("pos", "Vec3", "", "坐标", fill_style="stream"),
          FieldDef("items", "list,Item", "", "道具", fill_style="stream")]
    tp = os.path.join(ddir, "demo.xlsx")
    template.generate_data_xlsx(
        tm, TableDef(full_name="TbDemo", value_type="Demo", input="demo.xlsx"), tf, tp)
    wb2 = load_workbook(tp)
    ws2 = wb2.active
    merged = {str(r) for r in ws2.merged_cells.ranges}
    check({"C1:E1", "C2:E2", "C3:E3", "F1:I1"} <= merged,
          "多列字段的 ##var / ##type / ##group 各自合并成一个单元格：%s" % sorted(merged))
    check(str(ws2.cell(row=2, column=2).fill.fgColor.rgb).endswith("D9E7F7"),
          "单列字段仍是浅蓝表头")
    c1 = str(ws2.cell(row=1, column=3).fill.fgColor.rgb)
    c2 = str(ws2.cell(row=1, column=6).fill.fgColor.rgb)
    check(c1.endswith(template.MULTI_FILLS[0]) and c2.endswith(template.MULTI_FILLS[1])
          and c1 != c2,
          "多列字段填黄色且相邻两档深浅不同：%s / %s" % (c1[-6:], c2[-6:]))
    note_rows = [r for r in range(1, 12)
                 if str(ws2.cell(row=r, column=1).value or "") == "##"
                 and any(str(ws2.cell(row=r, column=c).value or "").find("[") >= 0
                         for c in range(2, (ws2.max_column or 1) + 1))]
    check(len(note_rows) == 1, "多列字段多出一行列序备注：%s" % note_rows)
    if note_rows:
        r0 = note_rows[0]
        cells = [ws2.cell(row=r0, column=c).value for c in range(1, (ws2.max_column or 1) + 1)]
        check(cells[2] == "pos[1/3]" and cells[4] == "pos[3/3]"
              and cells[8] == "items[4/4]",
              "备注写明每列是第几个元素：%s" % cells[1:])
    simple = os.path.join(ddir, "simple.xlsx")
    template.generate_data_xlsx(
        tm, TableDef(full_name="TbSimple", value_type="Demo", input="simple.xlsx"),
        [FieldDef("id", "int", "", "编号")], simple)
    wb3 = load_workbook(simple)
    ws3 = wb3.active
    check(not ws3.merged_cells.ranges,
          "没有多列字段时不做任何合并（不给普通表添乱）")
    wb2.close()
    wb3.close()
    shutil.rmtree(ddir, ignore_errors=True)

    print("== 18. 表专用自定义类型：独立定义文件 + 归属单选 ==")
    sdir = os.path.join(tempfile.gettempdir(), "lubankit_scope")
    shutil.rmtree(sdir, ignore_errors=True)
    ok, _ = projectwizard.create_project(sdir, ["xlsx"])
    check(ok, "建一个空工程")
    sm = projectconf.discover(sdir)
    exceldef.load_project(sm)
    sm.tables.append(TableDef(full_name="demo.TbDemo", value_type="Demo",
                              input="demo.xlsx", index="id"))
    exceldef.save_tables(sm)
    sm = projectconf.discover(sdir)
    exceldef.load_project(sm)
    check(sm.find_table("demo.TbDemo") is not None, "表已落盘")
    bp = exceldef.table_def_file(sm, "demo.TbDemo", "demo.xlsx")
    check(os.path.basename(bp) == "demo.xml",
          "表专用结构按模块名建文件：%s" % os.path.basename(bp))
    ok, msg = exceldef.save_bean(
        sm, BeanDef(full_name="demo.Extra", fields=[FieldDef("a", "int")]), bp)
    check(ok and os.path.exists(bp), "结构写进独立文件：" + str(msg).splitlines()[0])
    ep = exceldef.table_def_file(sm, "demo.TbDemo", "demo.xlsx")
    ok, msg = exceldef.save_enum(
        sm, EnumDef(full_name="demo.ExtraKind",
                    items=[EnumItem("A", "甲", "1"), EnumItem("B", "乙", "2")]), ep)
    check(ok and os.path.exists(ep), "枚举同样写进表专用 xml：%s" % os.path.basename(ep))
    conf2 = json.load(open(os.path.join(sdir, "luban.conf"), encoding="utf-8"))
    names = [os.path.basename(s["fileName"]) for s in conf2["schemaFiles"]]
    check("Defines" in names,
          "Defines 目录登记进了 luban.conf：%s" % names)
    check(not any("Extra" in str(x) for x in
                  (os.listdir(os.path.join(sdir, "Data")))), "没有多出来的杂物")
    sm2 = projectconf.discover(sdir)
    exceldef.load_project(sm2)
    got = {b.full_name: b.scope for b in sm2.beans}
    check(got.get("demo.Extra") == "demo.TbDemo",
          "重开后专用结构仍归本表：%s" % got)
    check({e.full_name: e.scope for e in sm2.enums}.get("demo.ExtraKind") == "demo.TbDemo",
          "重开后专用枚举仍归本表")
    check(not any(b.full_name == "demo.Extra" for b in sm2.beans if b._file.endswith("__beans__.xlsx")),
          "公共 __beans__ 里没有混进专用类型")
    pub = [b.full_name for b in sm2.beans if b._file.endswith("__beans__.xlsx")]
    check("demo.Extra" not in pub, "公共定义文件保持干净：%s" % pub)

    w4 = ui.MainWindow()
    w4.open_project(sdir)
    d4 = ui.CreateTypeDialog(w4, w4.model, "demo.TbDemo")
    check(d4.rb_table.isEnabled(), "有表名时「当前表专用」可选")
    d4.rb_table.setChecked(True)
    d4.ed_name.setText("Extra3")
    check(d4.table_scope() and d4.full_name() == "demo.Extra3",
          "归属当前表时类型名带模块前缀：%s" % d4.full_name())
    d4c = ui.CreateTypeDialog(w4, w4.model, None)
    check(not d4c.rb_table.isEnabled() and "表名" in d4c.rb_table.text(),
          "没有表名时给出可读提示而不是静默灰掉")
    w4._open_new_table(ask_leave=False)
    w4.nt_name.setText("demo.TbNew")
    check(w4._field_ctx(w4.nt_fields)[1] == "demo.TbNew",
          "新建表页把正在填的表名当作归属")
    d5 = ui.CreateTypeDialog(w4, w4.model, w4._field_ctx(w4.nt_fields)[1])
    check(d5.rb_table.isEnabled(), "新建表页填了表名后也能选「当前表专用」")

    # 界面链路：新建专用结构 -> 落到 Defines/demo.xml
    real_exec = ui.CreateTypeDialog.exec

    def _fake_exec(dlg):
        dlg.ed_name.setText("Extra4")
        dlg.rb_table.setChecked(True)
        return ui.QDialog.Accepted

    try:
        ui.CreateTypeDialog.exec = _fake_exec
        w4._load_table_editor("demo.TbDemo")
        w4._create_type_inline(0)
    finally:
        ui.CreateTypeDialog.exec = real_exec
    check(any(b.full_name == "demo.Extra4" for b in w4.model.beans),
          "界面新建的专用结构进了 model")
    check(os.path.exists(os.path.join(sdir, "Defines", "demo.xml")),
          "界面新建的专用结构落在 Defines/demo.xml")
    shutil.rmtree(sdir, ignore_errors=True)

    print("== 19. 左侧工程目录：双击文件用系统默认程序打开 ==")
    w5 = ui.MainWindow()
    w5.open_project(ROOT)
    it_file = None
    stack = [w5.tree.topLevelItem(i) for i in range(w5.tree.topLevelItemCount())]
    while stack and it_file is None:
        n = stack.pop(0)
        dd = n.data(0, ui.Qt.UserRole)
        if dd and dd[0] == "file":
            it_file = n
            break
        stack += [n.child(i) for i in range(n.childCount())]
    check(it_file is not None, "目录树里找得到一个文件节点")
    real_open_ext = w5._open_external
    opened = []
    w5._open_external = lambda p: opened.append(p)
    if it_file is not None:
        want = it_file.data(0, ui.Qt.UserRole)[1]
        w5.tree.itemDoubleClicked.emit(it_file, 0)
        check(opened[-1:] == [want],
              "双击文件 -> 交给系统默认程序：%s" % os.path.basename(want))
    opened.clear()
    w5.tree.itemDoubleClicked.emit(w5.tree.topLevelItem(0), 0)
    check(not opened, "双击目录节点不会去开外部程序（保持展开/收起）")
    w5._open_external = real_open_ext
    real_warn = ui.QMessageBox.warning
    ui.QMessageBox.warning = staticmethod(lambda *a, **k: None)
    try:
        w5._open_external(os.path.join(ROOT, "已经被删掉的.xlsx"))
        check(True, "文件已不存在时只提示，不抛异常")
    finally:
        ui.QMessageBox.warning = real_warn

    print("== 20. 管理自定义类型：引用统计 + 删除 ==")
    cdir = os.path.join(tempfile.gettempdir(), "lubankit_types")
    shutil.rmtree(cdir, ignore_errors=True)
    ok, _ = projectwizard.create_project(cdir, ["xlsx"])
    check(ok, "建一个空工程")
    cm = projectconf.discover(cdir)
    exceldef.load_project(cm)
    cm.tables.append(TableDef(full_name="demo.TbDemo", value_type="Reward",
                              input="demo.xlsx", index="id"))
    cm.tables.append(TableDef(full_name="demo.TbOther", value_type="Other",
                              input="other.xlsx", index="id"))
    exceldef.save_tables(cm)
    cm = projectconf.discover(cdir)
    exceldef.load_project(cm)
    r1 = exceldef.save_bean(cm, BeanDef(
        full_name="Reward", fields=[FieldDef("id", "int"), FieldDef("cost", "Cost")]))
    r2 = exceldef.save_bean(cm, BeanDef(
        full_name="Cost", fields=[FieldDef("amount", "int")]))
    r3 = exceldef.save_bean(
        cm, BeanDef(full_name="demo.Extra", fields=[FieldDef("a", "int")]),
        exceldef.table_def_file(cm, "demo.TbDemo", "bean"))
    r4 = exceldef.save_enum(cm, EnumDef(
        full_name="Quality", items=[EnumItem("A", "甲", "1")]))
    r5 = exceldef.save_enum(
        cm, EnumDef(full_name="demo.Kind", items=[EnumItem("A", "甲", "1")]),
        exceldef.table_def_file(cm, "demo.TbDemo", "enum"))
    check(all(x[0] for x in (r1, r2, r3, r4, r5)),
          "写入公共 / 表专用的结构与枚举")
    cm = projectconf.discover(cdir)
    exceldef.load_project(cm)

    check(ui.referenced_types("map,string,(Reward#sep=,)") == ["string", "Reward"],
          "容器里的元素类型也能解析出来")
    check(ui.referenced_types("int#ref=1-100") == ["int"], "标签不会被当成类型名")

    cfields = {
        "demo.TbDemo": [FieldDef("id", "int"), FieldDef("reward", "Reward"),
                        FieldDef("quality", "Quality"),
                        FieldDef("kinds", "list,demo.Kind")],
        "demo.TbOther": [FieldDef("id", "int"), FieldDef("kind", "demo.Kind")],
    }
    d6 = ui.ManageTypesDialog(w5, cm, cfields)
    names = [(e["kind"], e["name"], e["scope"]) for e in d6._entries]
    check(("bean", "demo.Extra", "demo.TbDemo") in names,
          "表专用结构归属本表：%s" % names)
    check(("enum", "Quality", "common") in names, "公共枚举在列表里")
    left = [d6.lst.item(i).text().strip() for i in range(d6.lst.count())]
    check(any(t.startswith("所有") for t in left)
          and any(t.startswith("公共") for t in left)
          and any(t.startswith("demo.TbDemo") for t in left)
          and any(t.startswith("demo.TbOther") for t in left),
          "左侧＝所有 + 公共 + 每张表名：%s" % left)

    def _entry(name):
        return next(e for e in d6._entries if e["name"] == name)

    check("demo.TbDemo.reward" in d6._ref_texts(_entry("Reward")),
          "公共类型显示「表名.字段名」：%s" % d6._ref_texts(_entry("Reward")))
    check(d6._ref_texts(_entry("Cost")) == ["公共.Reward.cost"],
          "公共结构里的字段引用也带归属名")
    check(d6._ref_texts(_entry("demo.Kind")) == ["kinds", "demo.TbOther.kind（跨表）"],
          "表专用类型默认只显示字段名，跨表引用补表名：%s"
          % d6._ref_texts(_entry("demo.Kind")))
    check(d6._ref_texts(_entry("demo.Extra")) == [], "没人用的类型没有引用位置")

    d6.lst.setCurrentRow(d6._index_of("all"))
    row = next(r for r in range(d6.tbl.rowCount())
               if "demo.Extra" in d6.tbl.item(r, 0).text())
    check(d6.tbl.item(row, 1).text() == "0"
          and d6.tbl.item(row, 1).foreground().color().name().upper() == "#B42318",
          "引用次数为 0 标红")
    d6.ck_unused.setChecked(True)
    check(d6.tbl.rowCount() == 1, "「只看没被引用的」筛出 1 个")
    d6.ck_unused.setChecked(False)

    real_q = ui.QMessageBox.exec
    real_info = ui.QMessageBox.information
    ui.QMessageBox.exec = lambda self: ui.QMessageBox.Yes
    ui.QMessageBox.information = staticmethod(lambda *a, **k: None)
    try:
        d6._del(_entry("demo.Extra"))
    finally:
        ui.QMessageBox.exec = real_q
        ui.QMessageBox.information = real_info
    cm3 = projectconf.discover(cdir)
    exceldef.load_project(cm3)
    check(cm3.find_bean("demo.Extra") is None, "删除落盘：demo.Extra 从定义文件消失")
    check(cm3.find_bean("Reward") is not None, "别的类型没被误删")
    check(not any(e["name"] == "demo.Extra" for e in d6._entries), "删完列表里也没了")

    r6 = exceldef.save_bean(cm3, BeanDef(full_name="demo.Child", parent="Reward",
                                        fields=[FieldDef("b", "int")]))
    check(r6[0], "加一个子结构继承 Reward")
    cm4 = projectconf.discover(cdir)
    exceldef.load_project(cm4)
    d7 = ui.ManageTypesDialog(w5, cm4, cfields)
    warned = []
    real_warn2 = ui.QMessageBox.warning
    ui.QMessageBox.warning = staticmethod(
        lambda *a, **k: warned.append(" ".join(str(x) for x in a[1:])))
    try:
        d7._del(next(e for e in d7._entries if e["name"] == "Reward"))
    finally:
        ui.QMessageBox.warning = real_warn2
    check(warned and "子结构" in warned[0], "有子类型时拒绝删除并说明原因：%s" % warned)
    cm5 = projectconf.discover(cdir)
    exceldef.load_project(cm5)
    check(cm5.find_bean("Reward") is not None, "被拒绝后 Reward 还在")
    shutil.rmtree(cdir, ignore_errors=True)

    print("== 21. 发布：覆盖写入 / 临时脚本 真跑一遍 ==")
    pdir = os.path.join(tempfile.gettempdir(), "lubankit_pub")
    shutil.rmtree(pdir, ignore_errors=True)
    ok, _ = projectwizard.create_project(pdir, ["xlsx"])
    check(ok, "建一个空工程")
    wn = ui.MainWindow()
    wn.open_project(pdir)
    dpub = ui.PublishDialog(wn, wn.model, wn.settings)
    dpub.show()
    dpub._fields["exe"].setText("lubankit_fake_luban.exe")   # 让脚本能跑但必然失败
    dpub._fields["codeTarget"].setEditText("cs-bin")
    dpub._fields["dataTarget"].setEditText("json")
    real_exec2 = ui.QMessageBox.exec
    ui.QMessageBox.exec = lambda self: ui.QMessageBox.Yes
    try:
        # ① 勾了覆盖：写回工程根目录的 gen.bat（先备份）
        dpub.ck_overwrite.setChecked(True)
        dpub._run()
    finally:
        ui.QMessageBox.exec = real_exec2
    bat = os.path.join(pdir, "gen.bat")
    check(os.path.exists(bat), "覆盖模式把 gen.bat 写回工程根目录")
    text = open(bat, encoding="gbk", errors="replace").read()
    check("-c cs-bin" in text and "-d json" in text and "--conf" in text,
          "写回的脚本带上界面里设置的参数")
    check("pause" in text and "if not defined" in text,
          "写回的脚本保留 pause，且工具运行时能跳过")
    baks = os.listdir(os.path.join(pdir, "_lubankit_backup"))
    check(any(f.startswith("gen_") and f.endswith(".bat") for f in baks),
          "覆盖前先备份原脚本：%s" % baks)
    # 脚本是怎么被执行的：直接调系统命令（Windows 就是 cmd /c）
    cmd0 = runner.script_command(os.path.join(pdir, "gen.bat"))
    check(len(cmd0) == 3 and os.path.basename(cmd0[0]).lower() in
          ("cmd.exe", "sh", "bash"), "用系统命令执行脚本：%s" % cmd0)
    check(cmd0[-1].endswith("gen.bat"), "系统命令的最后一个参数就是脚本路径")
    sc = publish.build_script(publish.blank_params(), "bat", pdir, True)
    check(("if not defined %s pause" % runner.PAUSE_ENV) in sc,
          "生成的脚本用与执行端同一个 pause 开关（名字不一致会卡死）")

    try:
        # ② 没勾覆盖：临时脚本建在**工程根目录**，跑失败则保留着方便排查
        dpub2 = ui.PublishDialog(wn, wn.model, wn.settings)
        dpub2.show()
        dpub2._fields["exe"].setText("lubankit_fake_luban.exe")
        dpub2._fields["codeTarget"].setEditText("cs-simple-json")
        dpub2.ck_overwrite.setChecked(False)
        ui.QMessageBox.exec = lambda self: ui.QMessageBox.Yes
        dpub2._run()
    finally:
        ui.QMessageBox.exec = real_exec2
    temps = [f for f in os.listdir(pdir)
             if f.startswith("gen_") and f.endswith(".bat")]
    check(len(temps) == 1, "临时脚本 gen_时间戳.bat 建在工程根目录：%s" % temps)
    check("-c cs-simple-json" in
          open(os.path.join(pdir, temps[0]), encoding="gbk",
               errors="replace").read(),
          "临时脚本里是这次的参数")
    text2 = open(bat, encoding="gbk", errors="replace").read()
    check("-c cs-bin" in text2 and "-c cs-simple-json" not in text2,
          "没勾覆盖时，工程里的 gen.bat 没被动过")
    os.remove(os.path.join(pdir, temps[0]))

    # ③ 导表成功：弹「是否删除临时脚本」，选删除就删、选保留就留着
    with open(os.path.join(pdir, "Luban.bat"), "w", encoding="gbk",
              newline="\r\n") as fh:
        fh.write("@echo off\r\necho fake-ok\r\nexit /b 0\r\n")
    for keep, want in ((ui.QMessageBox.No, True), (ui.QMessageBox.Yes, False)):
        dpub3 = ui.PublishDialog(wn, wn.model, wn.settings)
        dpub3.show()
        dpub3._fields["exe"].setText("Luban.bat")
        dpub3.ck_overwrite.setChecked(False)
        ui.QMessageBox.exec = lambda self, k=keep: k
        try:
            dpub3._run()
        finally:
            ui.QMessageBox.exec = real_exec2
        left = [f for f in os.listdir(pdir)
                if f.startswith("gen_") and f.endswith(".bat")]
        check(bool(left) == want,
              "成功后问是否删除临时脚本（%s）：剩下 %s"
              % ("保留" if want else "删除", left))
        for f in left:
            os.remove(os.path.join(pdir, f))
    shutil.rmtree(pdir, ignore_errors=True)

    print("== 22. 执行方式 / 工程目录自动查找 / 默认目标 ==")
    edir = os.path.join(tempfile.gettempdir(), "lubankit_pubexe")
    shutil.rmtree(edir, ignore_errors=True)
    projectwizard.create_project(edir, ["xlsx"])
    open(os.path.join(edir, "Luban.dll"), "wb").close()
    we = ui.MainWindow()
    we.open_project(edir)
    de = ui.PublishDialog(we, we.model, {})
    check(de._exe_mode() == "dll", "执行方式默认 dll：%s" % de._exe_mode())
    check(de._fields["exe"].text().replace("\\", "/") == "Luban.dll",
          "自动填上工程目录里的 Luban.dll：%s" % de._fields["exe"].text())
    check(de._fields["conf"].text() == "luban.conf",
          "配置文件取工程目录里的 luban.conf：%s" % de._fields["conf"].text())
    check(de._fields["codeTarget"].currentText() == "cs-simple-json",
          "代码目标默认 cs-simple-json：%s" % de._fields["codeTarget"].currentText())
    check(de._fields["dataTarget"].currentText() == "json",
          "数据格式默认 json：%s" % de._fields["dataTarget"].currentText())
    check(de.ed_cmd.text().startswith("dotnet"),
          "dll 方式命令前面补 dotnet：%s" % de.ed_cmd.text())
    open(os.path.join(edir, "Luban.exe"), "wb").close()
    de._fields["exeMode"].set_mode("exe")
    de._on_exe_mode("exe")
    check(de._exe_mode() == "exe", "可以切成 exe 方式")
    check(de._fields["exe"].text() == "Luban.exe",
          "切到 exe 后自动换成 Luban.exe：%s" % de._fields["exe"].text())
    check(not de.ed_cmd.text().startswith("dotnet"),
          "exe 方式不再补 dotnet：%s" % de.ed_cmd.text())
    de._fields["exeMode"].set_mode("dll")
    de._on_exe_mode("dll")
    check(de._fields["exe"].text() == "Luban.dll"
          and de.ed_cmd.text().startswith("dotnet"),
          "切回 dll 换回 Luban.dll 并补上 dotnet：%s" % de.ed_cmd.text())
    # 子目录里的 Luban 也能找到，并同步执行方式
    os.remove(os.path.join(edir, "Luban.dll"))
    os.remove(os.path.join(edir, "Luban.exe"))
    os.makedirs(os.path.join(edir, "Tools"), exist_ok=True)
    open(os.path.join(edir, "Tools", "Luban.exe"), "wb").close()
    de2 = ui.PublishDialog(we, we.model, {})
    check(de2._fields["exe"].text().replace("\\", "/") == "Tools/Luban.exe",
          "子目录里的 Luban 也能找到：%s" % de2._fields["exe"].text())
    check(de2._exe_mode() == "exe", "找到的是 exe，执行方式同步成 exe")
    we.close()
    shutil.rmtree(edir, ignore_errors=True)
    # 工程目录里没有 luban.conf / Luban → 留空，不瞎猜
    ndir3 = os.path.join(tempfile.gettempdir(), "lubankit_nofind")
    shutil.rmtree(ndir3, ignore_errors=True)
    os.makedirs(ndir3, exist_ok=True)
    wn3 = ui.MainWindow()
    wn3.open_project(ndir3)
    dn = ui.PublishDialog(wn3, wn3.model, {})
    check(dn._fields["conf"].text() == "",
          "工程目录里没有 luban.conf 就留空：%r" % dn._fields["conf"].text())
    check(dn._fields["exe"].text() == "",
          "没有 Luban.dll / Luban.exe 也留空：%r" % dn._fields["exe"].text())
    check(dn._exe_mode() == "dll", "什么都找不到时执行方式仍是 dll")
    wn3.close()
    shutil.rmtree(ndir3, ignore_errors=True)

    print("== 23. 表定义写在 Defines/*.xml 里：能读、能编辑、能写回 ==")
    import json as _json
    xdir = os.path.join(tempfile.gettempdir(), "lubankit_xml")
    shutil.rmtree(xdir, ignore_errors=True)
    os.makedirs(os.path.join(xdir, "Defines"))
    os.makedirs(os.path.join(xdir, "Data"))
    xpath = os.path.join(xdir, "Defines", "tables.xml")
    XML = """<?xml version="1.0" encoding="UTF-8"?>
<!-- 手写注释必须保住 -->
<module name="item">
  <!-- 注释里的 <table name="TbGhost" input="ghost.xlsx"/> 不算数 -->
  <table name="TbItem" value="Item" input="Data/item.xlsx" index="id" mode="map" group="c,s" comment="道具表"/>
  <table name="TbBag" value="Bag" input="bag.xlsx" mode="list" group="c" comment="背包表"/>
  <bean name="Item" comment="道具">
    <var name="id" type="int"/>
    <var name="name" type="string"/>
  </bean>
  <enum name="Kind" comment="类别">
    <var name="A" alias="甲" value="1"/>
  </enum>
</module>
"""
    with open(xpath, "w", encoding="utf-8") as f:
        f.write(XML)
    _json.dump({"schemaFiles": [{"fileName": "Defines", "type": ""}],
                "dataDir": "Data", "groups": [{"names": ["c", "s"]}],
                "targets": [{"name": "client"}]},
               open(os.path.join(xdir, "luban.conf"), "w", encoding="utf-8"),
               ensure_ascii=False, indent=2)

    xm = projectconf.discover(xdir)
    exceldef.load_project(xm)
    check(os.path.basename(xm.table_file) == "tables.xml" and
          [os.path.basename(p) for p in xm.table_files] == ["tables.xml"],
          "只声明 Defines 目录时也能定位到 xml 表定义：%s" % xm.table_files)
    names = [t.full_name for t in xm.tables]
    check(names == ["item.TbItem", "item.TbBag"],
          "xml 里的表按 module 前缀读出来：%s" % names)
    check("item.TbGhost" not in names, "注释里的 <table> 不当成真表")
    check(xm.find_bean("item.Item") is not None and
          [f.name for f in xm.find_bean("item.Item").fields] == ["id", "name"],
          "xml 里的 bean 也读到了")
    check(xm.find_enum("item.Kind") is not None, "xml 里的 enum 也读到了")

    # 数据文件能认出属于哪张表 -> 点开是编辑器，不是只读页
    wxml = ui.MainWindow()
    wxml.open_project(xdir)
    xdata = os.path.join(xdir, "Data", "item.xlsx")
    _wb = Workbook()
    _wb.active.append(["##var", "id", "name"])
    _wb.save(xdata)
    _wb.close()
    wxml._reindex_inputs()
    check(wxml._table_for_file(xdata) == "item.TbItem",
          "xml 里声明的表，其数据文件能匹配到表名（不再只读）：%s"
          % wxml._table_for_file(xdata))
    wxml._open_file(xdata)
    check(wxml.editor_stack.currentWidget() is wxml.page_preview,
          "点开该数据文件先进只读预览")
    wxml._on_edit_table()
    check(wxml.current == "item.TbItem", "编辑器当前表：%s" % wxml.current)

    # 写回：新增 / 修改 / 删除，注释与顺序保留
    xm.tables.append(TableDef(full_name="item.TbNew", value_type="New",
                              input="Data/new.xlsx", index="id", mode="map",
                              group="c,s", comment="新表"))
    for t in xm.tables:
        if t.full_name == "item.TbBag":
            t.comment = "背包表(改过)"
    xm.tables = [t for t in xm.tables if t.full_name != "item.TbBag"]
    ok, msg = exceldef.save_tables(xm)
    check(ok, "xml 表定义写回：" + str(msg)[:50])
    after = open(xpath, encoding="utf-8").read()
    check("手写注释必须保住" in after, "写回保留手写注释")
    check('name="TbNew"' in after, "新增的表写进 xml")
    check('name="TbBag"' not in after, "删掉的表从 xml 里移除")
    check("背包表(改过)" not in after and "TbBag" not in after, "删除是整行删掉")
    rows = [l for l in after.splitlines() if l.strip().startswith("<table ")]
    check(len(rows) == 2, "写回后 table 行数正确（注释里的不算）：%d" % len(rows))

    xm2 = projectconf.discover(xdir)
    exceldef.load_project(xm2)
    check([t.full_name for t in xm2.tables] == ["item.TbItem", "item.TbNew"],
          "重新读回：%s" % [t.full_name for t in xm2.tables])
    got = xm2.find_table("item.TbNew")
    check(got and (got.value_type, got.input, got.index, got.mode, got.comment)
          == ("New", "Data/new.xlsx", "id", "map", "新表"),
          "xml 表定义各属性往返一致")

    # 备份目录里的旧 xml 不能被当成定义文件（否则删掉的表又回来了）
    bak_files = [os.path.join(r, f) for r, _d, fs in
                 os.walk(os.path.join(xdir, "Defines", "_lubankit_backup")) for f in fs]
    check(bak_files, "写回前有备份：%s" % [os.path.basename(p) for p in bak_files])
    xm3 = projectconf.discover(xdir)
    check(all("_lubankit_backup" not in p for p in xm3.table_files),
          "备份目录里的 xml 不当成定义文件：%s" % xm3.table_files)
    exceldef.load_project(xm3)
    check("item.TbBag" not in [t.full_name for t in xm3.tables],
          "备份里的旧定义没有被读回来")

    # 解析失败（一张表都没读进来、也没记下来源）时不许把定义文件抹空
    xm3.tables = []
    xm3.loaded_tables = []
    ok3, _msg3 = exceldef.save_tables(xm3)
    check(not ok3 and "TbItem" in open(xpath, encoding="utf-8").read(),
          "疑似解析异常时拒绝清空定义文件：%s" % _msg3[:40])
    # 用户显式删光（这些表是本次读进来的）则放行
    xm4 = projectconf.discover(xdir)
    exceldef.load_project(xm4)
    xm4.tables = []
    ok4, _msg4 = exceldef.save_tables(xm4)
    left = [l for l in open(xpath, encoding="utf-8").read().splitlines()
            if l.strip().startswith("<table ")]
    check(ok4 and not left,
          "用户主动删光所有表时允许清空（注释里的示例行留着）：%s" % _msg4[:40])

    wxml.close()
    shutil.rmtree(xdir, ignore_errors=True)

    print("== 23b. xlsx 与 xml 混合的工程：各写各的文件 ==")
    mdir = os.path.join(tempfile.gettempdir(), "lubankit_mix")
    shutil.rmtree(mdir, ignore_errors=True)
    projectwizard.create_project(mdir, ["xlsx", "xml"])
    mm = projectconf.discover(mdir)
    exceldef.load_project(mm)
    mm.tables.append(TableDef(full_name="TbXlsx", value_type="Reward",
                              input="xlsx表.xlsx", index="id", mode="map", comment="xlsx 里的表"))
    ok, msg = exceldef.save_tables(mm)
    check(ok, "混合工程：先往 xlsx 写一张表：%s" % str(msg)[:40])
    ok, msg = xmldef.write_tables(os.path.join(mdir, "Defines", "tables.xml"), [
        TableDef(full_name="TbXml", value_type="Reward", input="xml表.xlsx",
                 index="id", mode="map", comment="xml 里的表")])
    check(ok, "混合工程：往 xml 写一张表：%s" % str(msg)[:40])

    mm2 = projectconf.discover(mdir)
    exceldef.load_project(mm2)
    files = {t.full_name: os.path.basename(t._file) for t in mm2.tables}
    check(files.get("TbXlsx") == "__tables__.xlsx", "xlsx 表记住来源：%s" % files)
    check(files.get("TbXml") == "tables.xml", "xml 表记住来源：%s" % files)
    check(len(mm2.table_files) == 2, "两个表定义文件都在：%s"
          % [os.path.basename(p) for p in mm2.table_files])

    for t in mm2.tables:
        t.comment = t.comment + "（改）"
    ok, msg = exceldef.save_tables(mm2)
    check(ok, "混合工程保存：%s" % str(msg)[:60])
    xml_text = open(os.path.join(mdir, "Defines", "tables.xml"), encoding="utf-8").read()
    check('name="TbXml"' in xml_text and "xml 里的表（改）" in xml_text,
          "xml 里的表写回了 xml")
    check('name="TbXlsx"' not in xml_text, "xlsx 里的表没有串到 xml")
    _ds2, _wb2, _w2 = exceldef.open_def_sheet(
        os.path.join(mdir, "Schemas", "__tables__.xlsx"), exceldef.TABLE_ALIASES)
    _map2, _ = exceldef.map_columns([c["l0"] for c in _ds2.columns],
                                    exceldef.TABLE_ALIASES)
    cmts = [_ds2.value(r, "comment", _map2) for r in _ds2.data_rows]
    check(any("xlsx 里的表（改）" in c for c in cmts), "xlsx 里的表写回了 xlsx：%s" % cmts)
    _wb2.close()
    shutil.rmtree(mdir, ignore_errors=True)

    print("== 24. 新建表：表声明文件类型（Excel Schema / XML Schema）==")
    ndir = os.path.join(tempfile.gettempdir(), "lubankit_deftype")
    shutil.rmtree(ndir, ignore_errors=True)
    projectwizard.create_project(ndir, ["xlsx"])
    w5 = ui.MainWindow()
    w5.open_project(ndir)
    check(w5.model.table_file.endswith(os.path.join("Schemas", "__tables__.xlsx")),
          "新建工程的主表定义在 Schemas 下：%s" % w5.model.table_file)

    # ① 默认 Excel Schema：声明进 __tables__.xlsx
    w5._open_new_table(ask_leave=False)
    w5.nt_name.setText("TbExcelOne")
    w5.nt_input.setText("excel_one.xlsx")
    check(w5.rb_nt_xlsx.isChecked(), "新建表默认 Excel Schema")
    w5._nt_create()
    t5 = w5.model.find_table("TbExcelOne")
    check(t5 is not None and os.path.basename(t5._file) == "__tables__.xlsx",
          "Excel Schema：声明指向 __tables__.xlsx")

    # ② 勾 XML Schema：没有 xml 表定义文件就新建一个并登记进 luban.conf
    w5._open_new_table(ask_leave=False)
    w5.nt_name.setText("TbXmlOne")
    w5.nt_input.setText("xml_one.xlsx")
    w5.rb_nt_xml.setChecked(True)
    w5._nt_create()
    xmlp = os.path.join(ndir, "Schemas", "tables.xml")
    t6 = w5.model.find_table("TbXmlOne")
    check(t6 is not None and os.path.basename(t6._file) == "tables.xml",
          "XML Schema：声明指向 tables.xml")
    check(os.path.exists(xmlp), "XML Schema：自动建出 tables.xml")
    _conf3 = json.load(open(os.path.join(ndir, "luban.conf"), encoding="utf-8"))
    check({"fileName": "Schemas/tables.xml", "type": "table"} in _conf3["schemaFiles"],
          "XML Schema：tables.xml 登记进 schemaFiles：%s" % _conf3["schemaFiles"])
    check(any(p.endswith("tables.xml") for p in w5.model.table_files),
          "XML Schema：新文件进了 table_files")

    w5._save_editor()
    _xmltext = open(xmlp, encoding="utf-8").read()
    check('name="TbXmlOne"' in _xmltext, "XML Schema：表声明真的写进了 xml")
    m6 = projectconf.discover(ndir)
    exceldef.load_project(m6)
    check(m6.find_table("TbExcelOne") is not None and m6.find_table("TbXmlOne") is not None,
          "重开工程两张表都在")
    check(os.path.basename(m6.find_table("TbXmlOne")._file) == "tables.xml",
          "重开后 xml 表仍记着来源文件")
    check(os.path.basename(m6.find_table("TbExcelOne")._file) == "__tables__.xlsx",
          "重开后 xlsx 表仍记着来源文件")
    # 再建一张 xml 表：应复用已有 tables.xml，而不是又开一个文件
    w5._open_new_table(ask_leave=False)
    w5.nt_name.setText("TbXmlTwo")
    w5.nt_input.setText("xml_two.xlsx")
    w5.rb_nt_xml.setChecked(True)
    w5._nt_create()
    check(os.path.basename(w5.model.find_table("TbXmlTwo")._file) == "tables.xml",
          "第二张 xml 表复用同一个 tables.xml")
    check(len([p for p in w5.model.table_files if p.endswith(".xml")]) == 1,
          "没有重复开 xml 定义文件：%s" % w5.model.table_files)
    shutil.rmtree(ndir, ignore_errors=True)

    print("== 25. 知识点窗口（顶部入口）==")
    w.act_knowledge()
    dk = w._knowledge_dlg
    check(dk is not None and dk.isVisible(), "点顶部「核心知识」能打开窗口")
    check([dk.tabs_k.tabText(i) for i in range(dk.tabs_k.count())] == ["核心知识", "常见问题"],
          "知识点列表分两栏")
    check(dk.lst_core.count() > 0, "核心知识列表有内容：%d 条" % dk.lst_core.count())
    check(dk.lst_faq.count() > 0, "常见问题列表有内容：%d 条" % dk.lst_faq.count())
    check(bool(dk.view.toPlainText().strip()), "正文区有内容")
    dk.ed_search.setText("多态")
    check(dk.lst_core.count() > 0, "窗口内搜索可用：核心知识 %d 条" % dk.lst_core.count())
    k0 = dk.lst_core.item(0).data(ui.Qt.UserRole) if dk.lst_core.count() else ""
    dk.focus_entry(k0)
    check(dk.ed_search.text() != "多态", "focus_entry 能跳到指定条目")
    dk.close()
    check(w._knowledge_dlg is None, "关掉窗口后主窗口的引用被清掉")

    print()
    # 测试结束：清空测试过程中写入的 QSettings（如发布测试写入的 luban_exe），
    # 别把测试用的「Luban.bat」等假路径当成历史记录污染真实环境。
    QSettings("LubanEditor", "LubanEditor").clear()
    if FAIL:
        print("FAILED: %d" % len(FAIL))
        for f in FAIL:
            print(" - " + f)
        return 1
    shutil.rmtree(ROOT, ignore_errors=True)
    print("ALL PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
