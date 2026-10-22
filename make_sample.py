# -*- coding: utf-8 -*-
"""生成一个示例工程 sample_project，用于：
1. 打包后自检（LubanEditor.exe --selftest sample_project）；
2. 发给别人时附带的"打开就能看懂"的演示工程。
"""
import json
import os
import shutil

from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "sample_project")

CONF = {
    "groups": [{"names": ["c"], "default": True},
               {"names": ["s"], "default": True},
               {"names": ["e"], "default": True}],
    "schemaFiles": [
        {"fileName": "Schemas/__tables__.xlsx", "type": "table"},
        {"fileName": "Defines", "type": ""},
    ],
    "dataDir": "Data",
    "targets": [{"name": "client", "manager": "Tables", "groups": ["c"], "topModule": "cfg"},
                {"name": "server", "manager": "Tables", "groups": ["s"], "topModule": "cfg"}],
}


def build(out=OUT):
    if os.path.exists(out):
        shutil.rmtree(out)
    d = os.path.join(out, "Data")          # 数据表
    s = os.path.join(out, "Schemas")       # Schema（表声明文件）
    df = os.path.join(out, "Defines")      # xml 定义（bean / enum）
    os.makedirs(d, exist_ok=True)
    os.makedirs(s, exist_ok=True)
    os.makedirs(df, exist_ok=True)
    with open(os.path.join(out, "luban.conf"), "w", encoding="utf-8") as f:
        json.dump(CONF, f, ensure_ascii=False, indent=2)

    # __tables__.xlsx
    wb = Workbook()
    ws = wb.active
    ws.title = "tables"
    ws.append(["##var", "full_name", "value_type", "read_schema_from_file", "input",
               "index", "mode", "group", "comment"])
    ws.append(["", "TbItem", "Item", "false", "item.xlsx", "id", "", "c,s", "道具表"])
    ws.append(["", "TbGlobal", "Global", "true", "global.xlsx", "", "one", "c,s", "全局配置"])
    wb.save(os.path.join(s, "__tables__.xlsx"))

    # Defines/common.xml（公共结构 + 枚举：Item 表结构 / Quality 枚举）
    with open(os.path.join(df, "common.xml"), "w", encoding="utf-8", newline="\n") as f:
        f.write(
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<!-- 类型定义（bean / enum），由 LubanEditor 创建 -->\n'
            '<module name="">\n'
            '  <bean name="Item" comment="道具">\n'
            '    <var name="id" type="int" comment="ID"/>\n'
            '    <var name="name" type="string" comment="名称"/>\n'
            '    <var name="quality" type="Quality" comment="品质"/>\n'
            '  </bean>\n'
            '  <enum name="Quality" comment="品质">\n'
            '    <var name="RED" alias="红" value="1"/>\n'
            '    <var name="BLUE" alias="蓝" value="2"/>\n'
            '  </enum>\n'
            '</module>\n')

    # 数据表 item.xlsx
    wb = Workbook()
    ws = wb.active
    ws.title = "item"
    ws.append(["##var", "id", "name", "quality"])
    ws.append(["##type", "int", "string", "Quality"])
    ws.append(["##group", "c,s", "c,s", "c,s"])
    ws.append(["##", "ID", "名称", "品质"])
    ws.append(["", 1001, "金币", "RED"])
    ws.append(["", 1002, "药水", "BLUE"])
    wb.save(os.path.join(d, "item.xlsx"))

    # 数据表 global.xlsx（纵表单例）
    wb = Workbook()
    ws = wb.active
    ws.title = "global"
    ws.append(["##column", "", "", ""])
    ws.append(["", "##type", "##", ""])
    ws.append(["bag_init_size", "int", "背包初始大小", 20])
    ws.append(["guild_open_level", "int", "公会开放等级", 10])
    wb.save(os.path.join(d, "global.xlsx"))
    return out


if __name__ == "__main__":
    print("示例工程已生成：", build())
